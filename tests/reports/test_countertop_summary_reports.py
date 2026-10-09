"""Countertop summaries repeat recorded result data without re-judging it."""

from __future__ import annotations

import re
from datetime import UTC, datetime
from fractions import Fraction
from io import BytesIO
from typing import TYPE_CHECKING
from uuid import UUID

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from app.schemas.visual_ui import (
    AgreementFactsOut,
    ArchitectComparedOut,
    ArchitectResultOut,
    CountertopPieceOut,
    CountertopResultOut,
    ExactValueOut,
    HoldOut,
    ReviewerDecisionOut,
    RowNotCheckedOut,
    WallLayoutOut,
)
from reports.findings_pdf import FindingsPdfInput, write_findings_pdf
from reports.spreadsheet import StoredFinding, write_stored_workbook
from units.imperial import format_inches
from verdict.outcomes import Outcome

if TYPE_CHECKING:
    # Imported late in the tests below, so each fails on its own assertion on main (#1093).
    from app.schemas.visual_ui import PageWithoutCountertopOut


def _exact(value: Fraction) -> ExactValueOut:
    return ExactValueOut(
        numerator=str(value.numerator),
        denominator=str(value.denominator),
        display=f'{format_inches(value)}"',
    )


def _result(
    label: str,
    outcome: Outcome | None,
    *,
    printed: Fraction | None,
    pieces: tuple[Fraction | None, ...],
    expected: Fraction | None,
    delta: Fraction | None,
    hold: HoldOut | None = None,
    decision: ReviewerDecisionOut | None = None,
    wall: WallLayoutOut | None = None,
) -> CountertopResultOut:
    return CountertopResultOut(
        finding_id=UUID(int=1) if outcome is not None else None,
        row_id=UUID(int=2 + len(label)),
        page_number=2,
        label=label,
        row_location=None,
        outcome=outcome,
        reviewer_decision=decision,
        needs_decision=outcome is None or outcome in (Outcome.FAIL, Outcome.REVIEW_REQUIRED),
        printed_overall=None if printed is None else _exact(printed),
        pieces=tuple(
            CountertopPieceOut(
                index=index,
                value=None if piece is None else _exact(piece),
                source="missing" if piece is None else "sealed",
                kind="cabinet",
            )
            for index, piece in enumerate(pieces)
        ),
        field_cut_per_end=_exact(Fraction(1)),
        field_cut_count=0,
        expected_total=None if expected is None else _exact(expected),
        delta=None if delta is None else _exact(delta),
        wall_layout=wall
        or WallLayoutOut(config="back_only", label="back wall only", source="reviewer"),
        agreement=AgreementFactsOut(
            both_readers_agreed_on_row=True,
            values_agreed=(True, *(True for _ in pieces)),
            code_clue_used=False,
        ),
        hold=hold,
    )


def _stored() -> StoredFinding:
    return StoredFinding(
        rule_id="CT-WIDTH-001",
        outcome="PASS",
        severity="CRITICAL",
        snapshot_id="snapshot",
        engine_version="test",
        trace={},
    )


def _pdf_text(data: bytes) -> str:
    return " ".join(
        " ".join(page.extract_text() or "" for page in PdfReader(BytesIO(data)).pages).split()
    )


def test_pdf_lists_checked_held_and_reviewed_countertops_before_findings() -> None:
    at = datetime(2026, 10, 9, 12, tzinfo=UTC)
    results = (
        _result(
            "Countertop row A",
            Outcome.PASS,
            printed=Fraction(24),
            pieces=(Fraction(10), Fraction(14)),
            expected=Fraction(24),
            delta=Fraction(0),
        ),
        _result(
            "Countertop row B",
            Outcome.FAIL,
            printed=Fraction(25),
            pieces=(Fraction(10), Fraction(14)),
            expected=Fraction(24),
            delta=Fraction(1),
            wall=WallLayoutOut(
                config="back_left_right", label="back wall and both ends", source="drawing clues"
            ),
        ),
        _result(
            "Countertop row C",
            None,
            printed=None,
            pieces=(Fraction(10), None),
            expected=None,
            delta=None,
            hold=HoldOut(
                code="row-partial", reason="A piece is missing; reviewer input is needed."
            ),
            decision=ReviewerDecisionOut(
                action="dismiss",
                note="Not checkable from the supplied drawing.",
                actor="Reviewer Example",
                time=at,
            ),
        ),
        _result(
            "Countertop between panels",
            Outcome.PASS,
            printed=Fraction(37),
            pieces=(Fraction(7), Fraction(23), Fraction(7)),
            expected=Fraction(37),
            delta=Fraction(0),
            wall=WallLayoutOut(
                config="back_only",
                label="back wall only; no field cut at the ends",
                source="between panels",
            ),
        ),
    )
    source = FindingsPdfInput(
        package_revision_id=UUID(int=11),
        revision_number=1,
        vendor=None,
        findings=(_stored(),),
        countertop_results=results,
    )

    text = _pdf_text(write_findings_pdf(source))

    assert text.index("COUNTERTOPS") < text.index("CT-WIDTH-001")
    assert "Countertop row A" in text and "PASS" in text
    assert '24"' in text and 'Difference 0"' in text
    assert "Countertop row B" in text and "FAIL" in text and 'Difference 1"' in text
    assert "back wall and both ends" in text and "drawing clues" in text
    assert "Countertop row C" in text and "A piece is missing" in text
    assert "Not checkable from the supplied drawing." in text
    assert "Reviewer Example" in text and "2026-10-09" in text
    assert "Countertop between panels" in text
    assert "between panels" in text and "no field cut at the ends" in text


def test_pdf_strip_uses_vector_drawing_for_a_result() -> None:
    from reports.findings_pdf import _draw_countertop_strip

    result = _result(
        "Countertop row A",
        Outcome.PASS,
        printed=Fraction(24),
        pieces=(Fraction(10), Fraction(14)),
        expected=Fraction(24),
        delta=Fraction(0),
    )

    class Recorder:
        def __init__(self) -> None:
            self.rectangles = 0
            self.lines = 0

        def rect(self, *_args: object, **_kwargs: object) -> None:
            self.rectangles += 1

        def line(self, *_args: object, **_kwargs: object) -> None:
            self.lines += 1

        def __getattr__(self, _name: str) -> object:
            return lambda *_args, **_kwargs: None

    canvas = Recorder()
    _draw_countertop_strip(canvas, result, x=0, y=0, width=420)

    assert canvas.rectangles >= len(result.pieces)
    assert canvas.lines >= 4


def test_pdf_countertop_card_keeps_strip_below_all_wrapped_details() -> None:
    from reports.findings_pdf import _countertop_strip_y

    top = 700.0
    detail_lines = 9
    last_detail_y = top - 31 - (detail_lines - 1) * 9
    strip_y = _countertop_strip_y(top, detail_lines)

    assert strip_y + 48 < last_detail_y
    assert strip_y - 25 > strip_y - 34


def test_workbook_has_exact_countertop_text_and_numeric_columns() -> None:
    result = _result(
        "Countertop row B",
        Outcome.FAIL,
        printed=Fraction(25),
        pieces=(Fraction(10), Fraction(14)),
        expected=Fraction(24),
        delta=Fraction(1),
    )

    book = load_workbook(BytesIO(write_stored_workbook((_stored(),), countertop_results=(result,))))

    assert "Countertops" in book.sheetnames
    sheet = book["Countertops"]
    headers = [cell.value for cell in sheet[1]]
    row = {header: sheet.cell(2, index + 1).value for index, header in enumerate(headers)}
    assert row["label"] == "Countertop row B"
    assert row["outcome"] == "FAIL"
    assert row["printed_overall"] == '25"'
    assert row["printed_overall_in"] == 25
    assert row["piece_1_in"] == 10
    assert row["piece_2_in"] == 14
    assert row["expected_total"] == '24"'
    assert row["difference"] == '1"'
    assert row["field_cut_per_end"] == '1"'
    assert row["field_cut_per_end_in"] == 1


# ---------------------------------------------------------------------------
# The architect line (#1054): "Matches the architect" beside each countertop
# ---------------------------------------------------------------------------


def _with_architect(result: CountertopResultOut, block: ArchitectResultOut) -> CountertopResultOut:
    return result.model_copy(update={"architect": block})


def _compared_overall(
    vendor: Fraction, architect: Fraction, outcome: Outcome
) -> ArchitectComparedOut:
    return ArchitectComparedOut(
        kind="overall",
        vendor_piece=None,
        vendor=_exact(vendor),
        architect=_exact(architect),
        delta=_exact(vendor - architect),
        vendor_display=_exact(vendor).display,
        architect_display=_exact(architect).display,
        delta_display=_exact(vendor - architect).display,
        outcome=outcome,
    )


def _checked_row() -> CountertopResultOut:
    return _result(
        "Countertop row A",
        Outcome.PASS,
        printed=Fraction(163, 4),
        pieces=(Fraction(10), Fraction(14)),
        expected=Fraction(163, 4),
        delta=Fraction(0),
    )


def test_pdf_card_says_whether_the_row_matches_the_architect() -> None:
    failed = _with_architect(
        _checked_row(),
        ArchitectResultOut(
            outcome=Outcome.FAIL,
            finding_id=UUID(int=40),
            reason="1 identifiers: 0 pass, 1 fail, 0 not found",
            needs_decision=True,
            compared=(_compared_overall(Fraction(163, 4), Fraction(44), Outcome.FAIL),),
            pairing_source="code",
        ),
    )
    not_compared = _with_architect(
        _result(
            "Countertop row B",
            Outcome.PASS,
            printed=Fraction(24),
            pieces=(Fraction(10), Fraction(14)),
            expected=Fraction(24),
            delta=Fraction(0),
        ),
        ArchitectResultOut(
            not_compared_reason="The architect prints nothing comparable for this row."
        ),
    )
    source = FindingsPdfInput(
        package_revision_id=UUID(int=12),
        revision_number=1,
        vendor=None,
        findings=(_stored(),),
        countertop_results=(failed, not_compared),
    )

    text = _pdf_text(write_findings_pdf(source))

    assert 'Matches the architect: FAIL (overall: vendor 40 3/4", architect 44")' in text
    assert (
        "Matches the architect: not compared: The architect prints nothing comparable for this row."
        in text
    )


def test_architect_line_for_a_pairing_waiting_for_the_reviewer() -> None:
    from reports.spreadsheet import architect_line

    waiting = _with_architect(
        _checked_row(),
        ArchitectResultOut(
            outcome=Outcome.REVIEW_REQUIRED,
            finding_id=UUID(int=41),
            reason="Pair the architect's dimension with the vendor's (one click).",
            needs_decision=True,
            pairing_source="both-ais",
        ),
    )

    assert architect_line(waiting) == (
        "Matches the architect: REVIEW_REQUIRED: Pair the architect's dimension with the vendor's "
        "(one click)."
    )


@pytest.mark.parametrize("source", ["code", "both-ais"])
def test_architect_line_says_to_confirm_a_one_judgment_result(source: str) -> None:
    from reports.spreadsheet import architect_line

    reason = (
        "Only code paired these; confirm that the architect's 3' - 8\" and the vendor's 40 3/4\" "
        "measure the same thing."
    )
    waiting = _with_architect(
        _checked_row(),
        ArchitectResultOut(
            outcome=Outcome.REVIEW_REQUIRED,
            finding_id=UUID(int=43),
            reason=reason,
            needs_decision=True,
            compared=(_compared_overall(Fraction(163, 4), Fraction(44), Outcome.REVIEW_REQUIRED),),
            pairing_source=source,
            pairing_judgments="code only" if source == "code" else "both AIs only",
        ),
    )

    assert architect_line(waiting) == (
        'Matches the architect: REVIEW_REQUIRED, to confirm (overall: vendor 40 3/4", '
        f'architect 44"): {reason}'
    )
    text = _pdf_text(
        write_findings_pdf(
            FindingsPdfInput(
                package_revision_id=UUID(int=12),
                revision_number=1,
                vendor=None,
                findings=(_stored(),),
                countertop_results=(waiting,),
            )
        )
    )
    assert "Matches the architect: REVIEW_REQUIRED, to confirm" in text


def test_architect_line_for_an_automatic_result_has_no_to_confirm() -> None:
    from reports.spreadsheet import architect_line

    automatic = _with_architect(
        _checked_row(),
        ArchitectResultOut(
            outcome=Outcome.FAIL,
            finding_id=UUID(int=44),
            needs_decision=True,
            compared=(_compared_overall(Fraction(163, 4), Fraction(44), Outcome.FAIL),),
            pairing_source="code+ais",
            pairing_judgments="code and both AIs",
        ),
    )

    assert architect_line(automatic) == (
        'Matches the architect: FAIL (overall: vendor 40 3/4", architect 44")'
    )


def test_workbook_names_the_pairing_judgments_and_says_to_confirm() -> None:
    one = _with_architect(
        _checked_row(),
        ArchitectResultOut(
            outcome=Outcome.REVIEW_REQUIRED,
            finding_id=UUID(int=45),
            needs_decision=True,
            compared=(_compared_overall(Fraction(163, 4), Fraction(44), Outcome.REVIEW_REQUIRED),),
            pairing_source="both-ais",
            pairing_judgments="both AIs only",
        ),
    )
    two = _with_architect(
        _checked_row(),
        ArchitectResultOut(
            outcome=Outcome.FAIL,
            finding_id=UUID(int=46),
            needs_decision=True,
            compared=(_compared_overall(Fraction(163, 4), Fraction(44), Outcome.FAIL),),
            pairing_source="code+ais",
            pairing_judgments="code and both AIs",
        ),
    )

    book = load_workbook(
        BytesIO(write_stored_workbook((_stored(),), countertop_results=(one, two)))
    )

    sheet = book["Countertops"]
    headers = [cell.value for cell in sheet[1]]
    rows = [
        {header: sheet.cell(index, column + 1).value for column, header in enumerate(headers)}
        for index in (2, 3)
    ]
    waiting = next(item for item in rows if item["architect_pairing_source"] == "both-ais")
    assert waiting["architect_pairing_judgments"] == "both AIs only — to confirm"
    assert waiting["architect_outcome"] == "REVIEW_REQUIRED"
    automatic = next(item for item in rows if item["architect_pairing_source"] == "code+ais")
    assert automatic["architect_pairing_judgments"] == "code and both AIs"
    assert headers.index("architect_pairing_judgments") == (
        headers.index("architect_pairing_source") + 1
    )


def test_the_strip_drawing_is_unchanged_by_the_architect_line() -> None:
    from reports.findings_pdf import _draw_countertop_strip

    class Recorder:
        def __init__(self) -> None:
            self.calls: list[tuple[str, tuple[object, ...]]] = []

        def __getattr__(self, name: str) -> object:
            return lambda *args, **_kwargs: self.calls.append((name, args))

    plain = Recorder()
    _draw_countertop_strip(plain, _checked_row(), x=0, y=0, width=420)
    with_line = Recorder()
    _draw_countertop_strip(
        with_line,
        _with_architect(
            _checked_row(),
            ArchitectResultOut(
                outcome=Outcome.FAIL,
                finding_id=UUID(int=40),
                compared=(_compared_overall(Fraction(163, 4), Fraction(44), Outcome.FAIL),),
            ),
        ),
        x=0,
        y=0,
        width=420,
    )

    assert with_line.calls == plain.calls


def test_workbook_has_the_architect_result_as_text_and_exact_numbers() -> None:
    result = _with_architect(
        _checked_row(),
        ArchitectResultOut(
            outcome=Outcome.FAIL,
            finding_id=UUID(int=40),
            needs_decision=True,
            compared=(_compared_overall(Fraction(163, 4), Fraction(44), Outcome.FAIL),),
            pairing_source="code",
        ),
    )
    third = _with_architect(
        _checked_row(),
        ArchitectResultOut(
            outcome=Outcome.PASS,
            finding_id=UUID(int=42),
            compared=(_compared_overall(Fraction(1, 3), Fraction(1, 3), Outcome.PASS),),
            pairing_source="reviewer",
        ),
    )

    book = load_workbook(
        BytesIO(write_stored_workbook((_stored(),), countertop_results=(result, third)))
    )

    sheet = book["Countertops"]
    headers = [cell.value for cell in sheet[1]]
    rows = [
        {header: sheet.cell(index, column + 1).value for column, header in enumerate(headers)}
        for index in (2, 3)
    ]
    row = next(item for item in rows if item["architect_pairing_source"] == "code")
    assert row["architect_outcome"] == "FAIL"
    assert (
        row["architect_comparison"]
        == 'overall: vendor 40 3/4", architect 44", difference -3 1/4" (FAIL)'
    )
    assert row["architect_not_compared_reason"] in ("", None)
    assert row["architect_overall_vendor_in"] == 40.75
    assert row["architect_overall_architect_in"] == 44
    assert row["architect_overall_difference_in"] == -3.25
    # A third is exact only as text: no rounded number is written.
    other = next(item for item in rows if item["architect_pairing_source"] == "reviewer")
    assert other["architect_overall_vendor_in"] in ("", None)
    assert (
        other["architect_comparison"]
        == 'overall: vendor 1/3", architect 1/3", difference 0" (PASS)'
    )
    # The piece columns still follow every named column.
    assert headers.index("piece_1_in") > headers.index("architect_overall_difference_in")


def test_workbook_says_why_a_row_was_not_compared() -> None:
    result = _with_architect(
        _checked_row(),
        ArchitectResultOut(not_compared_reason="No architect dimension is paired with this row."),
    )

    book = load_workbook(BytesIO(write_stored_workbook((_stored(),), countertop_results=(result,))))

    sheet = book["Countertops"]
    headers = [cell.value for cell in sheet[1]]
    row = {header: sheet.cell(2, index + 1).value for index, header in enumerate(headers)}
    assert row["architect_outcome"] == "NOT COMPARED"
    assert row["architect_not_compared_reason"] == "No architect dimension is paired with this row."


# ---------------------------------------------------------------------------
# A page whose countertop line the AIs did not agree on, and pages with none (#1093)
# ---------------------------------------------------------------------------

SPLIT_REASON = (
    "The two AIs did not agree on this page's countertop line (opus-5-5 picked line 2; "
    "sonnet-5-5 picked line 3), so nothing on it was read or checked. The reviewer decides "
    "this page."
)
NONE_REASON = "synthetic: this page shows only a wall elevation, no countertop dimension line"


def _split_item() -> CountertopResultOut:
    return _result(
        "Countertop row on page 3",
        Outcome.REVIEW_REQUIRED,
        printed=None,
        pieces=(),
        expected=None,
        delta=None,
        hold=HoldOut(code="row-choice-split", reason=SPLIT_REASON),
        decision=ReviewerDecisionOut(
            action="dismiss",
            note="TEST ONLY synthetic: the second line is the countertop",
            actor="Reviewer Example",
            time=datetime(2026, 10, 9, 12, tzinfo=UTC),
        ),
        wall=WallLayoutOut(config=None, label=None, source="not established"),
    ).model_copy(update={"page_number": 3, "field_cut_per_end": None, "field_cut_count": None})


def _pages_without_countertop() -> tuple[PageWithoutCountertopOut, ...]:
    from app.schemas.visual_ui import PageWithoutCountertopOut

    return (PageWithoutCountertopOut(page_number=4, reason=NONE_REASON),)


def test_pdf_lists_a_split_page_without_a_picture_and_the_pages_with_no_countertop() -> None:
    source = FindingsPdfInput(
        package_revision_id=UUID(int=12),
        revision_number=1,
        vendor=None,
        findings=(_stored(),),
        countertop_results=(_split_item(),),
        pages_without_countertop=_pages_without_countertop(),
    )

    text = _pdf_text(write_findings_pdf(source))

    assert "Countertop row on page 3" in text and "REVIEW_REQUIRED" in text
    assert "nothing on it was read or checked" in text
    assert "PAGES WITH NO COUNTERTOP FOUND" in text
    assert f"Page 4: {NONE_REASON}" in text
    assert text.index("PAGES WITH NO COUNTERTOP FOUND") < text.index("CT-WIDTH-001")


def test_pdf_lists_pages_with_no_countertop_even_with_no_countertop_item() -> None:
    source = FindingsPdfInput(
        package_revision_id=UUID(int=13),
        revision_number=1,
        vendor=None,
        findings=(_stored(),),
        pages_without_countertop=_pages_without_countertop(),
    )

    assert f"Page 4: {NONE_REASON}" in _pdf_text(write_findings_pdf(source))


def test_pdf_card_with_no_pieces_draws_no_countertop_picture(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from reports import findings_pdf

    drawn: list[str] = []
    monkeypatch.setattr(
        findings_pdf,
        "_draw_countertop_strip",
        lambda _canvas, result, **_kwargs: drawn.append(result.label),
    )
    held = _result(
        "Countertop row A",
        None,
        printed=None,
        pieces=(Fraction(10), None),
        expected=None,
        delta=None,
        hold=HoldOut(code="row-partial", reason="A piece is missing."),
    )
    write_findings_pdf(
        FindingsPdfInput(
            package_revision_id=UUID(int=14),
            revision_number=1,
            vendor=None,
            findings=(_stored(),),
            countertop_results=(held, _split_item()),
        )
    )

    assert drawn == ["Countertop row A"]


def test_workbook_has_the_split_page_and_the_pages_with_no_countertop() -> None:
    book = load_workbook(
        BytesIO(
            write_stored_workbook(
                (_stored(),),
                countertop_results=(_split_item(),),
                pages_without_countertop=_pages_without_countertop(),
            )
        )
    )

    sheet = book["Countertops"]
    headers = [cell.value for cell in sheet[1]]
    row = {header: sheet.cell(2, index + 1).value for index, header in enumerate(headers)}
    assert row["page"] == 3
    assert row["label"] == "Countertop row on page 3"
    assert row["outcome"] == "REVIEW_REQUIRED"
    assert row["pieces"] == "none recorded"
    assert row["hold_reason"] == SPLIT_REASON
    assert row["decision"] == "dismiss"

    empty = book["No Countertop Found"]
    assert [cell.value for cell in empty[1]] == ["page", "reason"]
    assert [cell.value for cell in empty[2]] == ["4", NONE_REASON]


_FOOTER = re.compile(r"GRANITI \+ NEXOLV - REVIEW FINDINGS PAGE (\d+)")


def _page_texts(data: bytes) -> list[str]:
    return [
        " ".join((page.extract_text() or "").split()) for page in PdfReader(BytesIO(data)).pages
    ]


def _many_results(count: int) -> tuple[CountertopResultOut, ...]:
    return tuple(
        _result(
            f"Countertop row {index}",
            Outcome.PASS,
            printed=Fraction(24),
            pieces=(Fraction(10), Fraction(14)),
            expected=Fraction(24),
            delta=Fraction(0),
        ).model_copy(update={"page_number": index + 1})
        for index in range(count)
    )


@pytest.mark.parametrize(
    ("results", "empty_pages", "changed"),
    [
        pytest.param(_many_results(1), False, False, id="one-countertop"),
        pytest.param(_many_results(1), False, True, id="after-project-values"),
        pytest.param((_split_item(),), True, False, id="split-and-no-countertop-list"),
        pytest.param((), True, False, id="only-no-countertop-list"),
        pytest.param(_many_results(9), True, False, id="countertops-continued"),
    ],
)
def test_pdf_countertop_section_has_no_empty_page_and_findings_start_their_own_page(
    results: tuple[CountertopResultOut, ...], empty_pages: bool, changed: bool
) -> None:
    from workflow.changed_values import UNAVAILABLE, ChangedValues

    source = FindingsPdfInput(
        package_revision_id=UUID(int=14),
        revision_number=1,
        vendor=None,
        findings=(_stored(),),
        countertop_results=results,
        changed_values=ChangedValues("unavailable", UNAVAILABLE, (), ()) if changed else None,
        pages_without_countertop=_pages_without_countertop() if empty_pages else (),
    )

    pages = _page_texts(write_findings_pdf(source))

    numbers = [int(match.group(1)) for page in pages for match in _FOOTER.finditer(page)]
    assert numbers == list(range(1, len(pages) + 1))
    assert all(_FOOTER.sub("", page).strip() for page in pages), "a page has only its footer"
    findings_start = next(index for index, page in enumerate(pages) if "CT-WIDTH-001" in page)
    assert pages[findings_start].startswith("FINDINGS REVISION 1")
    assert "COUNTERTOP" not in pages[findings_start]
    assert any("COUNTERTOPS" in page for page in pages[:findings_start])
    if len(results) > 4:
        assert any("COUNTERTOPS — CONTINUED" in page for page in pages[:findings_start])


# ---------------------------------------------------------------------------
# #1107: "drawn length not checked (no scale)" in the signed report
# ---------------------------------------------------------------------------

_NOT_CHECKED = "Drawn length not checked (no scale): piece 1, piece 2."


def _unwitnessed_row() -> CountertopResultOut:
    return _result(
        "Countertop row A",
        Outcome.REVIEW_REQUIRED,
        printed=Fraction(26),
        pieces=(Fraction(10), Fraction(14)),
        expected=None,
        delta=None,
    ).model_copy(update={"drawn_length_note": _NOT_CHECKED})


def test_pdf_card_says_drawn_length_not_checked_for_that_row() -> None:
    """Input: one row whose two pieces had no drawn-length witness, one row with every reading
    checked. Outcome: the signed PDF says so on the first card, once."""
    checked = _checked_row()
    source = FindingsPdfInput(
        package_revision_id=UUID(int=12),
        revision_number=1,
        vendor=None,
        findings=(_stored(),),
        countertop_results=(_unwitnessed_row(), checked),
    )

    text = _pdf_text(write_findings_pdf(source))

    assert text.count(_NOT_CHECKED) == 1


def test_workbook_has_a_drawn_length_column_after_the_named_ones() -> None:
    """Input: the same two rows. Outcome: the Countertops sheet's `drawn_length` column says it
    for the first row and is empty for the second; no earlier column moves and the per-piece
    columns still come last."""
    book = load_workbook(
        BytesIO(
            write_stored_workbook(
                (_stored(),), countertop_results=(_unwitnessed_row(), _checked_row())
            )
        )
    )

    sheet = book["Countertops"]
    headers = [cell.value for cell in sheet[1]]
    rows = [
        {header: sheet.cell(index, column + 1).value for column, header in enumerate(headers)}
        for index in (2, 3)
    ]
    by_outcome = {row["outcome"]: row for row in rows}
    assert by_outcome["REVIEW_REQUIRED"]["drawn_length"] == _NOT_CHECKED
    assert by_outcome["PASS"]["drawn_length"] in ("", None)
    assert headers.index("drawn_length") > headers.index("architect_overall_difference_in")
    assert headers.index("piece_1_in") > headers.index("drawn_length")


# #1108: a second countertop row an AI named on a page whose row was read is listed as not checked.
SECOND_ROW_REASON = (
    "An AI found a second countertop on this page (numbered line 3, named by opus-5-5). Only one "
    "countertop line per page is read, so it was not checked. Check it on the drawing."
)


def _rows_not_checked() -> tuple[RowNotCheckedOut, ...]:
    return (RowNotCheckedOut(page_number=2, reason=SECOND_ROW_REASON),)


def test_pdf_lists_a_second_countertop_row_that_was_not_checked() -> None:
    source = FindingsPdfInput(
        package_revision_id=UUID(int=15),
        revision_number=1,
        vendor=None,
        findings=(_stored(),),
        countertop_results=_many_results(1),
        pages_without_countertop=_pages_without_countertop(),
        rows_not_checked=_rows_not_checked(),
    )

    text = _pdf_text(write_findings_pdf(source))

    assert "SECOND COUNTERTOP ROWS NOT CHECKED" in text
    assert "Page 2: An AI found a second countertop on this page" in text
    assert "PAGES WITH NO COUNTERTOP FOUND" in text
    assert text.index("SECOND COUNTERTOP ROWS NOT CHECKED") < text.index("CT-WIDTH-001")


def test_pdf_lists_a_second_countertop_row_even_with_no_countertop_item() -> None:
    pages = _page_texts(
        write_findings_pdf(
            FindingsPdfInput(
                package_revision_id=UUID(int=16),
                revision_number=1,
                vendor=None,
                findings=(_stored(),),
                rows_not_checked=_rows_not_checked(),
            )
        )
    )

    assert any("SECOND COUNTERTOP ROWS NOT CHECKED" in page for page in pages)
    assert all(_FOOTER.sub("", page).strip() for page in pages), "a page has only its footer"


def test_pdf_refuses_rows_not_checked_of_another_type() -> None:
    with pytest.raises(TypeError, match="rows_not_checked"):
        FindingsPdfInput(
            package_revision_id=UUID(int=17),
            revision_number=1,
            vendor=None,
            findings=(_stored(),),
            rows_not_checked=_pages_without_countertop(),  # type: ignore[arg-type]
        )


def test_workbook_lists_a_second_countertop_row_that_was_not_checked() -> None:
    book = load_workbook(
        BytesIO(write_stored_workbook((_stored(),), rows_not_checked=_rows_not_checked()))
    )

    sheet = book["Second Rows Not Checked"]
    assert [cell.value for cell in sheet[1]] == ["page", "reason"]
    assert [cell.value for cell in sheet[2]] == ["2", SECOND_ROW_REASON]
    assert (
        "Second Rows Not Checked"
        not in load_workbook(BytesIO(write_stored_workbook((_stored(),)))).sheetnames
    )
