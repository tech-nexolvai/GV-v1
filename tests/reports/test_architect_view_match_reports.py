"""The signed reports name the architect view a countertop was compared with (#1168).

The architect's drawings as their own file: the workbook's architect line and the findings PDF (which
reuses the same line) end with `compared with <file>, page N, view X <title> (sheet S)`; a row waiting
for the reviewer's view choice says so; two new workbook columns carry the match state and the view.
A combined sheet (no match, no view) is exactly as before. Synthetic values only.
"""

from __future__ import annotations

from fractions import Fraction
from io import BytesIO
from uuid import UUID

from openpyxl import load_workbook

from app.schemas.architect_matches import ArchitectViewRefOut
from app.schemas.visual_ui import ArchitectMatchOut, ArchitectResultOut, CountertopResultOut
from reports.findings_pdf import FindingsPdfInput, write_findings_pdf
from reports.spreadsheet import COUNTERTOP_COLUMNS, architect_line, write_stored_workbook
from tests.reports.test_countertop_summary_reports import (
    _checked_row,
    _compared_overall,
    _pdf_text,
    _stored,
    _with_architect,
)
from verdict.outcomes import Outcome
from workflow.architect_row_plan import CHOOSE_ARCHITECT_VIEW

COMPARED_WITH = (
    "compared with the architect's drawings, page 2, view 3 SAMPLE ELEVATION (sheet Z-9)"
)
VIEW = ArchitectViewRefOut(
    view_id=UUID(int=300),
    document_id=UUID(int=301),
    document_version_id=UUID(int=302),
    file_name="the architect's drawings",
    page_number=2,
    sheet_number="Z-9",
    bubble="3",
    title="SAMPLE ELEVATION",
    scale_note=None,
    label="Page 2, view 3: SAMPLE ELEVATION (sheet Z-9)",
    region=None,
    picture_url=None,
    separated=True,
)


def _match(status: str, *, view: ArchitectViewRefOut | None = None) -> ArchitectMatchOut:
    return ArchitectMatchOut.model_validate(
        {
            "record_id": UUID(int=310),
            "status": status,
            "source": "automatic",
            "judgments": "code and both AIs" if view is not None else None,
            "code_verdict": None,
            "code_pick_view_id": None,
            "matched_view": view,
            "needs_decision": status == "needs_reviewer",
            "reason": None,
        }
    )


def _compared() -> CountertopResultOut:
    return _with_architect(
        _checked_row(),
        ArchitectResultOut(
            outcome=Outcome.FAIL,
            finding_id=UUID(int=320),
            compared=(_compared_overall(Fraction(48), Fraction(45), Outcome.FAIL),),
            pairing_source="code+ais",
            pairing_judgments="code and both AIs",
            match=_match("auto_matched", view=VIEW),
            compared_with=VIEW,
            compared_with_text=COMPARED_WITH,
        ),
    )


def _choose() -> CountertopResultOut:
    return _with_architect(
        _checked_row(),
        ArchitectResultOut(
            outcome=Outcome.REVIEW_REQUIRED,
            finding_id=UUID(int=321),
            reason=CHOOSE_ARCHITECT_VIEW,
            needs_decision=True,
            pairing_source="none",
            match=_match("needs_reviewer"),
        ),
    )


def _combined() -> CountertopResultOut:
    return _with_architect(
        _checked_row(),
        ArchitectResultOut(
            outcome=Outcome.FAIL,
            finding_id=UUID(int=322),
            compared=(_compared_overall(Fraction(48), Fraction(45), Outcome.FAIL),),
            pairing_source="code+ais",
            pairing_judgments="code and both AIs",
        ),
    )


def test_the_line_names_the_view_it_was_compared_with() -> None:
    assert architect_line(_compared()) == (
        f'Matches the architect: FAIL (overall: vendor 48", architect 45"); {COMPARED_WITH}'
    )


def test_a_row_waiting_for_the_view_choice_says_so() -> None:
    assert architect_line(_choose()) == (
        "Matches the architect: REVIEW_REQUIRED: Choose which of the architect's views shows this "
        "countertop (one click)."
    )


def test_a_combined_sheet_line_is_unchanged() -> None:
    assert architect_line(_combined()) == (
        'Matches the architect: FAIL (overall: vendor 48", architect 45")'
    )


def test_the_pdf_names_the_view_too() -> None:
    text = _pdf_text(
        write_findings_pdf(
            FindingsPdfInput(
                package_revision_id=UUID(int=12),
                revision_number=1,
                vendor=None,
                findings=(_stored(),),
                countertop_results=(_compared(), _choose()),
            )
        )
    )
    assert COMPARED_WITH in text
    assert "Choose which of the architect's views shows this countertop" in text


def test_the_workbook_has_the_match_state_and_the_view() -> None:
    book = load_workbook(
        BytesIO(
            write_stored_workbook(
                (_stored(),), countertop_results=(_compared(), _choose(), _combined())
            )
        )
    )
    sheet = book["Countertops"]
    headers = [cell.value for cell in sheet[1]]
    rows = [
        {header: sheet.cell(index, column + 1).value for column, header in enumerate(headers)}
        for index in (2, 3, 4)
    ]
    compared, choose, combined = rows
    assert compared["architect_match_status"] == "auto_matched"
    assert compared["architect_compared_with"] == COMPARED_WITH
    assert choose["architect_match_status"] == "needs_reviewer"
    assert choose["architect_compared_with"] in ("", None)
    # A combined sheet has no match: both new cells stay empty.
    assert combined["architect_match_status"] in ("", None)
    assert combined["architect_compared_with"] in ("", None)
    # Added at the end of the named columns: nothing before them moved; pieces still come last.
    assert COUNTERTOP_COLUMNS[-2:] == ("architect_match_status", "architect_compared_with")
    assert headers.index("architect_match_status") == headers.index("drawn_length") + 1
    assert headers.index("piece_1_in") > headers.index("architect_compared_with")


def test_a_pick_waiting_for_a_run_says_so() -> None:
    waiting = _choose().model_copy(
        update={
            "architect": _choose().architect.model_copy(
                update={
                    "match": _match("reviewer_confirmed", view=VIEW).model_copy(
                        update={"waits_for_run": True, "source": "reviewer"}
                    )
                }
            )
        }
    )
    assert architect_line(waiting) == (
        "Matches the architect: REVIEW_REQUIRED: Choose which of the architect's views shows this "
        "countertop (one click). (waits for the next check run)"
    )
