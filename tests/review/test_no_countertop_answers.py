"""A page every AI said has no countertop is never skipped silently (#1191).

Anant decided (2026-10-12, option c): an older reading that does not say what its "no line" meant is
blocking (`tests/api/test_unchosen_row_pages.py`), and every page still listed with no countertop
needs one reviewer click before sign-off: "Confirm: no countertop here", or "There is a countertop",
which makes the page a blocking item like a split page. The click is recorded with who and when,
carried only onto an identical page, frozen at sign-off and printed in the signed files.
Synthetic values only.
"""

from __future__ import annotations

import hashlib
import io
from collections.abc import Iterator
from decimal import Decimal
from pathlib import Path
from typing import Any
from uuid import UUID, uuid4

import pytest
from fastapi import HTTPException, Response
from fastapi.testclient import TestClient
from openpyxl import load_workbook  # type: ignore[import-untyped]
from pypdf import PdfReader
from sqlalchemy import Engine, func, select
from sqlalchemy.orm import Session

from alembic import command
from app.api.dependencies import get_artifact_store, get_session
from app.api.visual_countertops import (
    PAGE_NOT_LISTED,
    NoCountertopAnswerIn,
    _countertop_results_for_revision,
    answer_no_countertop_page,
)
from app.auth import Principal, Role, authenticate
from app.config import Settings
from app.db.session import session_factory
from app.main import create_app
from app.models import (
    Document,
    DocumentKind,
    DocumentVersion,
    NoCountertopAnswer,
    ObservationCandidate,
    PackageRevision,
    PackageRevisionDocument,
    PackageState,
    SourceArtifact,
)
from app.models.document import Page
from app.models.runs import ExtractionRun, TaskRun, WorkflowRun
from app.models.signed_exports import ApprovalExportSnapshot
from app.review.approval import (
    NoCountertopPagesUnconfirmed,
    approval_readiness,
    approval_readiness_many,
    approve_package,
)
from app.review.no_countertop import (
    pages_blocking_sign_off,
    reading_answer,
)
from app.review.session import open_session
from app.review.signed_exports import frozen_countertop_results, load_snapshot
from reports.findings_pdf import FindingsPdfInput, write_findings_pdf
from reports.spreadsheet import StoredFinding, write_stored_workbook
from storage.local import LocalStore
from tests.api.test_slot_rows import _run_current_checks
from tests.api.test_unchosen_row_pages import (
    NO_COUNTERTOP_KINDS,
    NONE_PICKS,
    NONE_WHY,
    _decide,
    _decided_rows,
    _live,
    _results,
    _revision_of,
    _unchosen_record,
    code_no_row_record,
)
from tests.app.postgres_fixture import alembic_config
from workflow.slot_row_scope import unchosen_page

pytest_plugins = ("tests.app.postgres_fixture",)

LISTED = NONE_PICKS + NO_COUNTERTOP_KINDS
REVIEWER = "synthetic-reviewer"
SIGNED_OFF = "This review is signed off. Start a new revision to change it."


@pytest.fixture
def session(postgres_engine: Engine) -> Iterator[Session]:
    config = alembic_config()
    config.attributes["database_url"] = postgres_engine.url.render_as_string(hide_password=False)
    command.upgrade(config, "head")
    opened = session_factory(postgres_engine)()
    try:
        yield opened
    finally:
        opened.close()


def _reviewer(
    project_id: UUID, *, roles: frozenset[Role] = frozenset({Role.REVIEWER})
) -> Principal:
    return Principal(id=REVIEWER, roles=roles, projects=frozenset({project_id}))


def _answer(
    session: Session, project_id: UUID, package_id: UUID, record_id: UUID, answer: str
) -> Any:
    principal = _reviewer(project_id)
    return answer_no_countertop_page(
        principal,
        principal,
        session,
        project_id,
        package_id,
        record_id,
        NoCountertopAnswerIn.model_validate({"answer": answer}),
        Response(),
    )


def _client(session: Session, principal: Principal, root: Path) -> TestClient:
    app = create_app(Settings(database_url="postgresql+psycopg://unused/unused"))
    app.dependency_overrides[authenticate] = lambda: principal
    app.dependency_overrides[get_session] = lambda: session
    store = LocalStore(root=root / "signed", ticket_secret=b"synthetic-test")
    app.dependency_overrides[get_artifact_store] = lambda: store
    return TestClient(app, raise_server_exceptions=False)


def _listed_page(session: Session, tmp_path: Path) -> tuple[UUID, UUID, PackageRevision, UUID]:
    """Two checked and decided countertops (pages 1 and 2) and page 3, which both AIs said has no
    countertop (with kinds, so it is listed). The review waits for sign-off."""
    project_id, package_id = _decided_rows(session)
    revision = _revision_of(session, package_id)
    record_id = _unchosen_record(session, revision.id, page_index=2, flags=LISTED, reason=NONE_WHY)
    _run_current_checks(session, package_id, tmp_path)
    _decide(session, revision.id)
    revision.state = PackageState.AWAITING_REVIEW.value
    session.commit()
    return project_id, package_id, revision, record_id


def _review_session(session: Session, revision_id: UUID) -> UUID:
    review = open_session(session, package_revision_id=revision_id, reviewer=REVIEWER)
    session.commit()
    return review.id


# --- sign-off waits for every listed page ---------------------------------------------------------


def test_sign_off_is_refused_until_every_listed_page_is_confirmed(
    session: Session, tmp_path: Path
) -> None:
    project_id, package_id, revision, record_id = _listed_page(session, tmp_path)

    # Nothing is pre-selected: the page is listed with no confirmation, and sign-off waits.
    (listed,) = _results(session, package_id).pages_without_countertop
    assert (listed.page_number, listed.record_id, listed.confirmation) == (3, record_id, None)
    readiness = approval_readiness(session, revision.id)
    assert readiness.blocking_findings == 0
    assert readiness.unconfirmed_no_countertop_pages == (3,)
    assert not readiness.can_approve
    assert readiness.reason == (
        "Confirm the pages listed with no countertop before signing off: page 3 is not confirmed "
        'yet. On each, choose "Confirm: no countertop here" or "There is a countertop".'
    )
    # The Documents summary reads the same rule.
    assert approval_readiness_many(session, [revision.id])[revision.id] == readiness

    review_id = _review_session(session, revision.id)
    refused = _client(session, _reviewer(project_id), tmp_path).post(
        f"/api/v1/projects/{project_id}/review-sessions/{review_id}/approve"
    )
    assert refused.status_code == 409
    assert "page 3 is not confirmed yet" in refused.json()["message"]
    with pytest.raises(NoCountertopPagesUnconfirmed, match="page 3 is not confirmed yet"):
        approve_package(session, principal=_reviewer(project_id), review_session_id=review_id)
    session.rollback()

    answered = _answer(session, project_id, package_id, record_id, "no_countertop")

    (confirmed,) = answered.pages_without_countertop
    assert confirmed.confirmation is not None
    assert confirmed.confirmation.confirmed_by == REVIEWER
    assert confirmed.confirmation.from_revision is None
    stored = session.scalars(select(NoCountertopAnswer)).one()
    assert (stored.answer, stored.answered_by, stored.page_number) == (
        "no_countertop",
        REVIEWER,
        3,
    )
    assert stored.created_at == confirmed.confirmation.confirmed_at
    readiness = approval_readiness(session, revision.id)
    assert readiness.unconfirmed_no_countertop_pages == ()
    assert readiness.can_approve, readiness.reason
    decision = approve_package(
        session, principal=_reviewer(project_id), review_session_id=review_id
    )
    session.commit()
    assert decision.approval.package_revision_id == revision.id


def test_there_is_a_countertop_makes_the_page_a_blocking_split_item(
    session: Session, tmp_path: Path
) -> None:
    project_id, package_id, revision, record_id = _listed_page(session, tmp_path)

    answered = _answer(session, project_id, package_id, record_id, "countertop")

    # No longer listed: it is an item, held like a split page, with nothing read on it yet.
    assert answered.pages_without_countertop == ()
    (item,) = [item for item in answered.items if item.row_id == record_id]
    assert item.hold is not None and item.hold.code == "row-choice-split"
    # The hold names who said there is a countertop, and when (it is printed in the signed files).
    said = item.hold.reason
    assert said.startswith(f"{REVIEWER} said on ") and "that this page has a countertop" in said
    assert item.finding_id is None and item.needs_decision
    readiness = approval_readiness(session, revision.id)
    assert readiness.unchecked_countertop_pages == (3,)
    assert readiness.reason == (
        "Page 3 has no check result yet. Run the checks before signing off."
    )
    # The page is no longer listed, so a second answer is refused in plain words.
    with pytest.raises(HTTPException) as refused:
        _answer(session, project_id, package_id, record_id, "no_countertop")
    assert refused.value.status_code == 409 and refused.value.detail == PAGE_NOT_LISTED

    # The next check run makes it one "needs you" result, decided like any split page.
    _run_current_checks(session, package_id, tmp_path)
    (finding,) = [f for f in _live(session, revision.id) if f.scope_row_candidate_id == record_id]
    assert finding.outcome == "REVIEW_REQUIRED"
    assert finding.reason == said
    readiness = approval_readiness(session, revision.id)
    assert readiness.blocking_finding_ids == (finding.id,)
    assert readiness.unchecked_countertop_pages == ()
    _decide(session, revision.id)
    session.commit()
    assert approval_readiness(session, revision.id).can_approve


def test_there_is_a_countertop_after_the_last_run_asks_for_a_check_run(
    session: Session, tmp_path: Path
) -> None:
    """The answer is an input the checks read, like a wall answer (#1137): a run older than it does
    not count."""
    from app.review.approval import _revisions_with_unchecked_inputs

    project_id, package_id, revision, record_id = _listed_page(session, tmp_path)
    _answer(session, project_id, package_id, record_id, "no_countertop")
    assert revision.id not in _revisions_with_unchecked_inputs(session, [revision.id])

    _answer(session, project_id, package_id, record_id, "countertop")
    assert revision.id in _revisions_with_unchecked_inputs(session, [revision.id])


def _older_page(session: Session, tmp_path: Path) -> tuple[UUID, UUID, PackageRevision, UUID]:
    """Like `_listed_page`, but page 3 is an older reading: both AIs picked 0 with no kind."""
    project_id, package_id = _decided_rows(session)
    revision = _revision_of(session, package_id)
    record_id = _unchosen_record(
        session, revision.id, page_index=2, flags=NONE_PICKS, reason=NONE_WHY
    )
    _run_current_checks(session, package_id, tmp_path)
    _decide(session, revision.id)
    revision.state = PackageState.AWAITING_REVIEW.value
    session.commit()
    return project_id, package_id, revision, record_id


def test_an_older_reading_blocks_until_confirmed_and_then_is_a_confirmed_no_countertop_page(
    session: Session, tmp_path: Path
) -> None:
    """Review finding 9: an older all-0 reading gets the same two buttons. Until answered it blocks
    sign-off; "Confirm" makes it a confirmed no-countertop page, never a "not checkable" result."""
    project_id, package_id, revision, record_id = _older_page(session, tmp_path)
    # The check stage wrote no "needs you" result for it.
    assert all(f.scope_row_candidate_id != record_id for f in _live(session, revision.id))
    readiness = approval_readiness(session, revision.id)
    assert readiness.unconfirmed_no_countertop_pages == (3,) and not readiness.can_approve
    review_id = _review_session(session, revision.id)
    with pytest.raises(NoCountertopPagesUnconfirmed, match="page 3 is not confirmed yet"):
        approve_package(session, principal=_reviewer(project_id), review_session_id=review_id)
    session.rollback()

    answered = _answer(session, project_id, package_id, record_id, "no_countertop")

    (page,) = answered.pages_without_countertop
    assert page.older_reading and page.confirmation is not None
    assert page.confirmation.confirmed_by == REVIEWER
    # A run after the answer still writes nothing for the page, and nothing waits for a run.
    _run_current_checks(session, package_id, tmp_path)
    assert all(f.scope_row_candidate_id != record_id for f in _live(session, revision.id))
    _decide(session, revision.id)
    session.commit()
    assert approval_readiness(session, revision.id).can_approve
    decision = approve_package(
        session, principal=_reviewer(project_id), review_session_id=review_id
    )
    session.commit()
    assert decision.approval.package_revision_id == revision.id


def test_an_older_reading_with_there_is_a_countertop_takes_the_split_path(
    session: Session, tmp_path: Path
) -> None:
    project_id, package_id, revision, record_id = _older_page(session, tmp_path)

    answered = _answer(session, project_id, package_id, record_id, "countertop")

    assert answered.pages_without_countertop == ()
    (item,) = [item for item in answered.items if item.row_id == record_id]
    assert item.hold is not None and item.hold.code == "row-choice-split"
    _run_current_checks(session, package_id, tmp_path)
    (finding,) = [f for f in _live(session, revision.id) if f.scope_row_candidate_id == record_id]
    assert finding.outcome == "REVIEW_REQUIRED"


def test_a_package_for_another_product_never_waits_for_a_countertop_result(
    session: Session, tmp_path: Path
) -> None:
    """`countertops_checked=False`: a cabinet package's checks never write a countertop result, so
    a page a reviewer said has a countertop does not wait for one."""
    from app.models import Package
    from app.review.no_countertop import blocking_pages, countertops_checked

    project_id, package_id, revision, record_id = _listed_page(session, tmp_path)
    _answer(session, project_id, package_id, record_id, "countertop")
    assert pages_blocking_sign_off(session, [revision.id])[revision.id].unchecked == (3,)
    package = session.get(Package, package_id)
    assert package is not None
    package.product_type = "cabinet"
    session.commit()

    assert not countertops_checked("cabinet") and countertops_checked(None)
    assert pages_blocking_sign_off(session, [revision.id])[revision.id].unchecked == ()
    assert approval_readiness(session, revision.id).unchecked_countertop_pages == ()
    assert _results(session, package_id).items  # the page is still shown as an item
    assert blocking_pages((), {}, set(), countertops_checked=False).unchecked == ()


# --- repeat clicks --------------------------------------------------------------------------------


def test_a_repeated_confirm_records_nothing_and_keeps_the_first_person_and_time(
    session: Session, tmp_path: Path
) -> None:
    project_id, package_id, _revision, record_id = _listed_page(session, tmp_path)
    url = (
        f"/api/v1/projects/{project_id}/packages/{package_id}/no-countertop-pages/"
        f"{record_id}/answer"
    )
    first = _client(session, _reviewer(project_id), tmp_path).post(
        url, json={"answer": "no_countertop"}
    )
    other = Principal(
        id="synthetic-second", roles=frozenset({Role.REVIEWER}), projects=frozenset({project_id})
    )
    again = _client(session, other, tmp_path).post(url, json={"answer": "no_countertop"})

    assert (first.status_code, again.status_code) == (201, 200)
    assert session.scalar(select(func.count()).select_from(NoCountertopAnswer)) == 1
    confirmed = again.json()["pages_without_countertop"][0]["confirmation"]
    assert confirmed == first.json()["pages_without_countertop"][0]["confirmation"]
    assert confirmed["confirmed_by"] == REVIEWER


def test_answers_on_one_revision_are_taken_one_at_a_time(session: Session, tmp_path: Path) -> None:
    """The newest answer is read under a transaction-scoped advisory lock, before anything is
    written, so a "Confirm" sent with a "There is a countertop" waits for it and is then refused
    (the page is no longer listed) instead of overriding it."""
    from sqlalchemy import event

    project_id, package_id, _revision, record_id = _listed_page(session, tmp_path)
    statements: list[str] = []

    def seen(*args: object) -> None:
        statements.append(str(args[2]))

    engine = session.get_bind()
    event.listen(engine, "before_cursor_execute", seen)
    try:
        _answer(session, project_id, package_id, record_id, "countertop")
    finally:
        event.remove(engine, "before_cursor_execute", seen)
    lock = next(i for i, text in enumerate(statements) if "pg_advisory_xact_lock" in text)
    insert = next(
        i for i, text in enumerate(statements) if "INSERT INTO no_countertop_answers" in text
    )
    reads = [i for i, text in enumerate(statements) if "FROM no_countertop_answers" in text]
    assert lock < min(reads) < insert
    # The later "Confirm" finds the page no longer listed.
    with pytest.raises(HTTPException) as refused:
        _answer(session, project_id, package_id, record_id, "no_countertop")
    assert refused.value.status_code == 409


# --- carried only onto an identical page ----------------------------------------------------------


def _reread(session: Session, revision_id: UUID) -> UUID:
    """A newer slot-reader run on the revision: only its records count from now on."""
    workflow_run = WorkflowRun(package_revision_id=revision_id, engine_run_id=str(uuid4()))
    session.add(workflow_run)
    session.flush()
    task = TaskRun(
        workflow_run_id=workflow_run.id,
        idempotency_key=str(uuid4()),
        task_type="extract",
        attempt=1,
        outcome="SUCCEEDED",
    )
    session.add(task)
    session.flush()
    run = ExtractionRun(
        task_run_id=task.id,
        extractor="extraction.form_reader",
        extractor_version="slot-reader-v1-test",
        config_hash="synthetic-test",
        dpi=150,
    )
    session.add(run)
    session.flush()
    return run.id


def _record_on(session: Session, run_id: UUID, page: Page, flags: tuple[str, ...]) -> UUID:
    record = ObservationCandidate(
        document_version_id=page.document_version_id,
        page_id=page.id,
        extraction_run_id=run_id,
        raw_text="",
        polygon=[[0, 0], [612, 0], [612, 792], [0, 792]],
        coordinate_space="image",
        ambiguity_flags=["slot-reader-row-choice", "row-choice:0", *flags],
        review_reason=NONE_WHY,
    )
    session.add(record)
    session.flush()
    return record.id


@pytest.mark.parametrize(
    ("flags", "carried"),
    [
        (LISTED, True),
        # The same page answered by other readers: another reading answer, asked again.
        (
            (
                "row-pick:opus-6-0:0",
                "row-pick:sonnet-6-0:0",
                "row-kind:opus-6-0:no_countertop",
                "row-kind:sonnet-6-0:no_countertop",
            ),
            False,
        ),
    ],
    ids=["identical-page-and-answer", "answer-changed"],
)
def test_a_confirmation_carries_over_a_rereading_only_with_the_same_answer(
    session: Session, tmp_path: Path, flags: tuple[str, ...], carried: bool
) -> None:
    project_id, package_id, revision, record_id = _listed_page(session, tmp_path)
    _answer(session, project_id, package_id, record_id, "no_countertop")
    page = session.get(Page, session.get(ObservationCandidate, record_id).page_id)  # type: ignore[union-attr]
    assert page is not None

    again = _record_on(session, _reread(session, revision.id), page, flags)
    session.commit()

    (listed,) = _results(session, package_id).pages_without_countertop
    assert listed.record_id == again
    assert (listed.confirmation is not None) is carried
    assert pages_blocking_sign_off(session, [revision.id])[revision.id].unconfirmed == (
        () if carried else (3,)
    )


def _two_revision_pages(
    session: Session, hashes: tuple[tuple[str, str], ...]
) -> tuple[PackageRevision, PackageRevision, list[Page], list[Page]]:
    """Revision 1 and 2 of one package, each with its own issue of one shop drawing whose pages
    have these content hashes (revision 1's, revision 2's)."""
    from tests.api.test_packages import _new_package

    package_id = _new_package(session)
    one = session.scalars(
        select(PackageRevision).where(PackageRevision.package_id == package_id)
    ).one()
    document = Document(package_id=package_id, kind=DocumentKind.SHOP.value)
    session.add(document)
    session.flush()
    two = PackageRevision(
        package_id=package_id, revision_number=2, state="UPLOADED", supersedes_id=one.id
    )
    session.add(two)
    session.flush()
    pages: list[list[Page]] = []
    for position, revision in enumerate((one, two)):
        digest = hashlib.sha256(f"synthetic issue {position}".encode()).hexdigest()
        artifact = SourceArtifact(storage_key=f"synthetic/{uuid4()}", sha256=digest, size=1)
        session.add(artifact)
        session.flush()
        version = DocumentVersion(
            document_id=document.id,
            source_artifact_id=artifact.id,
            sha256=digest,
            page_count=len(hashes),
        )
        session.add(version)
        session.flush()
        session.add(
            PackageRevisionDocument(
                package_revision_id=revision.id,
                package_id=package_id,
                document_id=document.id,
                document_version_id=version.id,
            )
        )
        issue = [
            Page(
                document_version_id=version.id,
                index=index,
                content_hash=pair[position],
                width_pt=Decimal(612),
                height_pt=Decimal(792),
                rotation=0,
                has_vector_text=True,
            )
            for index, pair in enumerate(hashes)
        ]
        session.add_all(issue)
        session.flush()
        pages.append(issue)
    return one, two, pages[0], pages[1]


def test_revision_two_keeps_a_confirmation_only_on_an_identical_page_with_a_known_hash(
    session: Session,
) -> None:
    empty = hashlib.sha256(b"").hexdigest()
    one, two, old_pages, new_pages = _two_revision_pages(
        session,
        (
            ("1" * 64, "1" * 64),  # page 1: byte-identical on revision 2
            ("2" * 64, "3" * 64),  # page 2: changed on revision 2
            (empty, empty),  # page 3: content unknown on both
        ),
    )
    run_one = _reread(session, one.id)
    for page in old_pages:
        record = session.get(ObservationCandidate, _record_on(session, run_one, page, LISTED))
        assert record is not None
        listed = unchosen_page(page.index, record, page.content_hash)
        session.add(
            NoCountertopAnswer(
                package_revision_id=one.id,
                row_choice_candidate_id=record.id,
                page_id=page.id,
                page_number=page.index + 1,
                page_content_hash=page.content_hash,
                reading_answer=reading_answer(listed),
                answer="no_countertop",
                answered_by=REVIEWER,
            )
        )
    run_two = _reread(session, two.id)
    for page in new_pages:
        _record_on(session, run_two, page, LISTED)
    session.commit()

    results = _countertop_results_for_revision(session, two.package_id, two)
    confirmations = {
        page.page_number: page.confirmation for page in results.pages_without_countertop
    }
    assert set(confirmations) == {1, 2, 3}
    carried = confirmations[1]
    assert carried is not None
    assert (carried.from_revision, carried.confirmed_by) == (1, REVIEWER)
    assert confirmations[2] is None  # the page changed
    assert confirmations[3] is None  # an unknown page is never "the same page"
    assert pages_blocking_sign_off(session, [two.id])[two.id].unconfirmed == (2, 3)
    # Revision 1 shows its own answers, given there, with no "from revision".
    earlier = _countertop_results_for_revision(session, one.package_id, one)
    assert [
        (page.page_number, page.confirmation is not None and page.confirmation.from_revision)
        for page in earlier.pages_without_countertop
    ] == [(1, None), (2, None), (3, None)]


def test_an_answer_given_on_revision_two_never_reaches_revision_one(session: Session) -> None:
    one, two, old_pages, new_pages = _two_revision_pages(session, (("1" * 64, "1" * 64),))
    _record_on(session, _reread(session, one.id), old_pages[0], LISTED)
    record = session.get(
        ObservationCandidate, _record_on(session, _reread(session, two.id), new_pages[0], LISTED)
    )
    assert record is not None
    session.add(
        NoCountertopAnswer(
            package_revision_id=two.id,
            row_choice_candidate_id=record.id,
            page_id=new_pages[0].id,
            page_number=1,
            page_content_hash="1" * 64,
            reading_answer=reading_answer(unchosen_page(0, record, "1" * 64)),
            answer="no_countertop",
            answered_by=REVIEWER,
        )
    )
    session.commit()

    assert pages_blocking_sign_off(session, [one.id, two.id]) == {
        one.id: pages_blocking_sign_off(session, [one.id])[one.id],
        two.id: pages_blocking_sign_off(session, [two.id])[two.id],
    }
    assert pages_blocking_sign_off(session, [one.id])[one.id].unconfirmed == (1,)
    assert pages_blocking_sign_off(session, [two.id])[two.id].unconfirmed == ()


# --- who may answer, and when ---------------------------------------------------------------------


@pytest.mark.parametrize(
    "roles",
    [frozenset({Role.RULE_ADMIN}), frozenset()],
    ids=["rule-admin", "no-role"],
)
def test_only_a_reviewer_may_answer(
    session: Session, tmp_path: Path, roles: frozenset[Role]
) -> None:
    project_id, package_id, _revision, record_id = _listed_page(session, tmp_path)
    url = (
        f"/api/v1/projects/{project_id}/packages/{package_id}/no-countertop-pages/"
        f"{record_id}/answer"
    )

    refused = _client(session, _reviewer(project_id, roles=roles), tmp_path).post(
        url, json={"answer": "no_countertop"}
    )
    accepted = _client(session, _reviewer(project_id), tmp_path).post(
        url, json={"answer": "no_countertop"}
    )

    # A refused action answers like an absence (`app/auth/dependencies._refuse`): 404, never 403.
    assert refused.status_code == 404, refused.text
    assert accepted.status_code == 201, accepted.text
    assert (
        accepted.json()["pages_without_countertop"][0]["confirmation"]["confirmed_by"] == REVIEWER
    )
    assert session.scalar(select(func.count()).select_from(NoCountertopAnswer)) == 1


def test_an_answer_must_be_one_of_the_two_buttons(session: Session, tmp_path: Path) -> None:
    project_id, package_id, _revision, record_id = _listed_page(session, tmp_path)
    url = (
        f"/api/v1/projects/{project_id}/packages/{package_id}/no-countertop-pages/"
        f"{record_id}/answer"
    )
    client = _client(session, _reviewer(project_id), tmp_path)

    assert client.post(url, json={}).status_code == 422
    assert client.post(url, json={"answer": "maybe"}).status_code == 422
    assert session.scalar(select(func.count()).select_from(NoCountertopAnswer)) == 0


def test_answers_are_refused_after_sign_off(session: Session, tmp_path: Path) -> None:
    project_id, package_id, revision, record_id = _listed_page(session, tmp_path)
    _answer(session, project_id, package_id, record_id, "no_countertop")
    approve_package(
        session,
        principal=_reviewer(project_id),
        review_session_id=_review_session(session, revision.id),
    )
    session.commit()

    refused = _client(session, _reviewer(project_id), tmp_path).post(
        f"/api/v1/projects/{project_id}/packages/{package_id}/no-countertop-pages/"
        f"{record_id}/answer",
        json={"answer": "countertop"},
    )

    assert refused.status_code == 409
    assert refused.json()["message"] == SIGNED_OFF
    assert session.scalar(select(func.count()).select_from(NoCountertopAnswer)) == 1


# --- frozen at sign-off and printed ---------------------------------------------------------------


def test_the_confirmation_is_frozen_at_sign_off_and_printed_in_the_signed_files(
    session: Session, tmp_path: Path
) -> None:
    project_id, package_id, revision, record_id = _listed_page(session, tmp_path)
    _answer(session, project_id, package_id, record_id, "no_countertop")
    decision = approve_package(
        session,
        principal=_reviewer(project_id),
        review_session_id=_review_session(session, revision.id),
    )
    session.commit()

    snapshot = session.scalars(
        select(ApprovalExportSnapshot).where(
            ApprovalExportSnapshot.approval_id == decision.approval.id
        )
    ).one()
    frozen = frozen_countertop_results(session, load_snapshot(session, snapshot))
    (page,) = frozen.pages_without_countertop
    assert page.record_id == record_id
    assert page.confirmation is not None and page.confirmation.confirmed_by == REVIEWER
    stamp = page.confirmation.confirmed_at.isoformat()

    finding = StoredFinding(
        rule_id="CT-WIDTH-001",
        outcome="PASS",
        severity="FLAG",
        snapshot_id="synthetic-snapshot",
        engine_version="0.1.0",
        trace={},
        reason="Synthetic.",
    )
    pdf = write_findings_pdf(
        FindingsPdfInput(
            package_revision_id=revision.id,
            revision_number=1,
            vendor="Synthetic vendor",
            findings=(finding,),
            countertop_results=frozen.items,
            pages_without_countertop=frozen.pages_without_countertop,
        )
    )
    text = " ".join(
        " ".join((sheet.extract_text() or "").split()) for sheet in PdfReader(io.BytesIO(pdf)).pages
    )
    assert f"Confirmed no countertop: page 3 — {REVIEWER}," in text
    book = load_workbook(
        io.BytesIO(
            write_stored_workbook(
                (finding,), pages_without_countertop=frozen.pages_without_countertop
            )
        )
    )
    sheet = book["No Countertop Found"]
    assert [cell.value for cell in sheet[1]] == ["page", "reason", "confirmed"]
    assert [cell.value for cell in sheet[2]] == [
        "3",
        NONE_WHY,
        f"Confirmed no countertop: page 3 — {REVIEWER}, {stamp}",
    ]


# --- carried to revision 2 ------------------------------------------------------------------------


def _carried_to_revision_two(
    session: Session, answer: str
) -> tuple[PackageRevision, PackageRevision, UUID]:
    """Revision 1 answered page 1 (`answer`); revision 2's page 1 is byte-identical and was read to
    the same answer. Returns both revisions and revision 2's record."""
    one, two, old_pages, new_pages = _two_revision_pages(session, (("1" * 64, "1" * 64),))
    record = session.get(
        ObservationCandidate, _record_on(session, _reread(session, one.id), old_pages[0], LISTED)
    )
    assert record is not None
    session.add(
        NoCountertopAnswer(
            package_revision_id=one.id,
            row_choice_candidate_id=record.id,
            page_id=old_pages[0].id,
            page_number=1,
            page_content_hash="1" * 64,
            reading_answer=reading_answer(unchosen_page(0, record, "1" * 64)),
            answer=answer,
            answered_by=REVIEWER,
        )
    )
    new_record = _record_on(session, _reread(session, two.id), new_pages[0], LISTED)
    one.state = PackageState.SUPERSEDED.value
    session.commit()
    return one, two, new_record


def test_a_confirmation_carried_to_revision_two_needs_the_signers_still_applies_tick(
    session: Session,
) -> None:
    """Review fix (C5 reversed): like #1180's carried wall answers, a confirmation given on revision
    1's identical page is listed at sign-off, and the signer must say it still applies."""
    from app.review.approval import CarriedAnswersUnconfirmed
    from tests.api.test_revisions import _scoped_finding

    _one, two, _record = _carried_to_revision_two(session, "no_countertop")
    pages = pages_blocking_sign_off(session, [two.id])[two.id]
    assert pages.unconfirmed == ()
    assert pages.carried == ((1, f"Page 1: no countertop confirmed on revision 1 ({REVIEWER})"),)
    # A result to sign, as a checked revision 2 would have.
    _scoped_finding(session, two, _record)
    two.state = PackageState.AWAITING_REVIEW.value
    session.commit()
    review_id = _review_session(session, two.id)
    signer = Principal(
        id="synthetic-signer", roles=frozenset({Role.ADMIN}), projects=frozenset({two.package_id})
    )

    with pytest.raises(
        CarriedAnswersUnconfirmed, match=r"Page 1: no countertop confirmed on revision 1"
    ):
        approve_package(session, principal=signer, review_session_id=review_id)
    session.rollback()

    decision = approve_package(
        session, principal=signer, review_session_id=review_id, carried_answers_confirmed=True
    )
    session.commit()
    assert decision.approval.package_revision_id == two.id


def test_a_carried_there_is_a_countertop_needs_no_tick_and_gets_a_result_in_the_check_stage(
    session: Session, tmp_path: Path
) -> None:
    from tests.workflow.test_stages import _publish_rulebook
    from workflow.stages import DatabaseStages

    _publish_rulebook(session)
    _one, two, record_id = _carried_to_revision_two(session, "countertop")
    pages = pages_blocking_sign_off(session, [two.id])[two.id]
    assert (pages.unconfirmed, pages.unchecked, pages.carried) == ((), (1,), ())
    (item,) = _countertop_results_for_revision(session, two.package_id, two).items
    assert (
        item.hold is not None
        and "on revision 1 that this page has a countertop" in item.hold.reason
    )

    two.state = PackageState.NEEDS_INPUT.value
    session.commit()
    DatabaseStages(
        store=LocalStore(root=tmp_path / "store", ticket_secret=b"synthetic-test")
    ).run_checks(session, two.id)

    (finding,) = [f for f in _live(session, two.id) if f.scope_row_candidate_id == record_id]
    assert finding.outcome == "REVIEW_REQUIRED"
    assert pages_blocking_sign_off(session, [two.id])[two.id].unchecked == ()


def test_there_is_a_countertop_on_a_page_code_read_nothing_on_says_code_found_no_row(
    session: Session, tmp_path: Path
) -> None:
    """#1193: a page code read no row on (no vendor drawing it could read) is answered with the same
    two buttons; "There is a countertop" makes it a blocking item whose reason says code, not the
    AIs, found no row."""
    project_id, package_id = _decided_rows(session)
    revision = _revision_of(session, package_id)
    record_id = code_no_row_record(session, revision.id, page_index=2)
    _run_current_checks(session, package_id, tmp_path)
    _decide(session, revision.id)
    revision.state = PackageState.AWAITING_REVIEW.value
    session.commit()

    answered = _answer(session, project_id, package_id, record_id, "countertop")

    assert answered.pages_without_countertop == ()
    (item,) = [item for item in answered.items if item.row_id == record_id]
    assert item.hold is not None and item.hold.code == "row-choice-split"
    assert "although code found no countertop row to read on it" in item.hold.reason
    assert "both AIs" not in item.hold.reason
    _run_current_checks(session, package_id, tmp_path)
    (finding,) = [f for f in _live(session, revision.id) if f.scope_row_candidate_id == record_id]
    assert finding.outcome == "REVIEW_REQUIRED" and finding.reason == item.hold.reason
