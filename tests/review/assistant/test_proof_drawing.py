"""The zero-click red/green proof drawing (#1581), on synthetic PDFs, $0.

- which items code clouds by itself (only Wrong items code proved), and where the rest go;
- the colours: red and green legible on white and in greyscale print;
- the words: the sheet note `46"  should be  37"` (red, grey, green) and one proof sentence per mark;
- the rendered PDF: cover counts (Wrong / Right / Not compared), stamp, assumptions, the
  sheet's note, the proof entry, the page labels; the reviewer's decisions as overrides; what could
  not be compared is never on the drawing (#1619): it is the separate Not-compared list;
- the plan check never waits: code's likely answers, written on the cover.
"""

from __future__ import annotations

import io
import json
from pathlib import Path
from typing import Any

import numpy as np
import pikepdf
import pypdfium2 as pdfium  # type: ignore[import-untyped]
import pytest
from pypdf import PdfReader
from reportlab.lib.colors import Color  # type: ignore[import-untyped]
from reportlab.pdfgen import canvas  # type: ignore[import-untyped]

from app.review.assistant.proof import (
    AI_ONLY_REASON,
    PAIRING_REASON,
    RULE_LINE,
    code_proved,
    governs,
    needs_pairing,
    proof_counts,
    proof_group,
    proof_line,
    proof_of,
    should_be,
    unproved_reason,
    vendor_own,
    vendor_value,
)
from app.review.assistant.proof_drawing import (
    GREEN,
    RED,
    STAMP_WORDS,
    contrast,
    proof_markup_pdf,
    stamp_line,
)

WIDTH, HEIGHT = 1000.0, 700.0
from tests.review.assistant.reading_fixture import two_routes

#: The full check's devil's advocate passing an item (#1611): the pairing is proved.
ADVOCATE_PASSED: dict[str, Any] = {
    "result": "passed",
    "answers": {"spans": "agree", "extent": "same_extent"},
    "note": "",
    "model": "devil's advocate",
}
#: Where the sheet prints its values (shown points, origin top-left: x, baseline y).
VENDOR_46 = (600.0, 520.0)
ARCHITECT_37 = (200.0, 200.0)
VENDOR_12 = (400.0, 600.0)
VENDOR_30 = (800.0, 300.0)


def _sheet(path: Path, pages: int = 1) -> Path:
    """A sheet: the architect's 3'-1" and the vendor's 46", 12" and 30" printed as text."""
    out = io.BytesIO()
    c = canvas.Canvas(out, pagesize=(WIDTH, HEIGHT), invariant=1)
    for _ in range(pages):
        c.setFont("Helvetica", 10)
        c.drawString(ARCHITECT_37[0], HEIGHT - ARCHITECT_37[1], "3'-1\"")
        c.drawString(VENDOR_46[0], HEIGHT - VENDOR_46[1], '46"')
        c.drawString(VENDOR_12[0], HEIGHT - VENDOR_12[1], '12"')
        c.drawString(VENDOR_30[0], HEIGHT - VENDOR_30[1], '30"')
        c.showPage()
    c.save()
    path.write_bytes(out.getvalue())
    return path


def _box(at: tuple[float, float], page: int = 1, file: str = "sheet.pdf") -> dict[str, Any]:
    x, y = at
    return {"page": page, "x0": x - 2, "y0": y - 10, "x1": x + 22, "y1": y + 3, "file": file}


def _checks(**changes: Any) -> dict[str, Any]:
    checks: dict[str, Any] = {
        "printed_architect": "match",
        "printed_vendor": "match",
        "second_read": "agrees",
        "rule": None,
        "arithmetic": None,
        "lead_checked": True,
        "ledger": None,
        "units": None,
        "adjudication": None,
        "critic": None,
    }
    checks.update(changes)
    return checks


def _proved(**changes: Any) -> dict[str, Any]:
    """A Wrong item code proved: both values printed at their spots, read again by a model."""
    finding: dict[str, Any] = {
        "id": "F1",
        "page": 1,
        "vs_page": None,
        "item": "Buffet outlet height",
        "category": "cabinet",
        "architect": "3'-1\"",
        "vendor": '46"',
        "status": "different",
        "source": "architect",
        "why": "outlet 9 in higher",
        "trust_level": "confirmed",
        "witnesses": ["printed"],
        "checks": _checks(),
        "inches": {"architect": "37", "vendor": "46", "difference": "9"},
        "vendor_box": _box(VENDOR_46),
        "arch_box": _box(ARCHITECT_37),
        # The full check's devil's advocate proved both numbers measure the same thing (#1611).
        "critic": ADVOCATE_PASSED,
    }
    finding.update(changes)
    return two_routes(finding)


def _ai_only(**changes: Any) -> dict[str, Any]:
    return (
        _proved(
            id="F2",
            item="Base cabinet height",
            architect='34"',
            vendor='12"',
            trust_level="to_check",
            witnesses=[],
            checks=_checks(printed_architect="no_text", printed_vendor="no_text"),
            inches={"architect": "34", "vendor": "12", "difference": "-22"},
            vendor_box=_box(VENDOR_12),
            arch_box=None,
        )
        | changes
    )


def _decision(**changes: Any) -> dict[str, Any]:
    return (
        _proved(
            id="F3",
            item="Counter overall width",
            architect="no width printed",
            vendor='30"',
            status="check",
            trust_level="to_check",
            witnesses=[],
            keep_reason="architect_prints_nothing",
            question="The architect prints no size: is the vendor's 30\" right?",
            vendor_box=_box(VENDOR_30),
            arch_box=None,
        )
        | changes
    )


def _right(**changes: Any) -> dict[str, Any]:
    return (
        _proved(
            id="F4",
            item="Toe kick height",
            architect='4"',
            vendor='4"',
            status="same",
            compared=True,
            vendor_box=None,
            arch_box=None,
        )
        | changes
    )


# ---- which items code clouds by itself ---------------------------------------------------------


@pytest.mark.parametrize(
    ("finding", "group"),
    [
        (_proved(), "clouded"),
        (_proved(trust_level="likely", checks=_checks(second_read="not_run")), "clouded"),
        (_proved(witnesses=["rule"], source="rules", rule_check_id="C1"), "clouded"),
        # What code could not prove or compare is Not compared (#1619), never a question.
        (_ai_only(), "not_compared"),
        # Two AI readings agreeing are Likely, but no code witness: never clouded alone.
        (_ai_only(trust_level="likely"), "not_compared"),
        # A check that disagrees makes it "AI reading only".
        (_proved(trust_level="to_check"), "not_compared"),
        (_decision(), "not_compared"),
        (_right(), "right"),
        (_right(compared=False), "not_compared"),
        # The reviewer's decision overrides, both ways.
        (_ai_only(decision={"decision": "agree", "by": "R"}), "clouded"),
        (_decision(decision={"decision": "agree", "by": "R"}), "clouded"),
        (_proved(decision={"decision": "not_a_problem", "by": "R"}), "dismissed"),
    ],
)
def test_each_item_goes_to_its_group_by_trust_and_decision(
    finding: dict[str, Any], group: str
) -> None:
    assert proof_group(finding) == group


def test_a_scaled_value_or_a_vendor_side_only_witness_is_never_clouded() -> None:
    # The printed scales are wrong wherever GV resized a pasted drawing: never proof.
    assert not code_proved(_proved(architect='no size printed; scales about 74"'))
    assert not code_proved(_proved(check_reason="scaled"))
    # Millimetres and inches agreeing back the vendor's reading only: with an architect chain
    # code cannot read as one size, code does not know what it should be.
    chain = _proved(
        architect='1\'-0", 1\'-6", 8" wall to wall (no overall printed)',
        witnesses=["units"],
        inches={"architect": None, "vendor": "107.75", "difference": None},
    )
    assert not code_proved(chain)
    assert code_proved({**chain, "inches": {"architect": "106", "vendor": "107.75"}})


def test_an_architect_comparison_needs_a_pairing_proof_a_rule_or_sum_does_not() -> None:
    """Reading the vendor's number right (its mm and inches agreeing, or printed at its spot)
    does not prove the architect's number measures the same thing (a counter's underside against
    its top). The one signal is the full check's (#1611): code paired the two itself, or the
    devil's advocate passed the item. Without it the item is grey, with the reason."""
    unpaired = _proved(witnesses=["units"], critic=None)
    assert proof_group(unpaired) == "not_compared"
    assert proof_of(unpaired)["not_compared"] == "different_reference"
    assert unproved_reason(unpaired) == PAIRING_REASON
    assert unproved_reason(_ai_only()) == AI_ONLY_REASON
    # The advocate objecting: grey with its own reason.
    objects = {
        **ADVOCATE_PASSED,
        "result": "objects",
        "note": "Not proved: a blind second pick chose p23.D12, not p23.D10.",
    }
    objected = _proved(critic=objects, trust_level="to_check")
    assert proof_group(objected) == "not_compared"
    assert unproved_reason(objected).startswith("Not proved: a blind second pick chose p23.D12")
    # Another model's critic passing is not the advocate's proof.
    other = _proved(critic={**ADVOCATE_PASSED, "model": "critic"})
    assert proof_group(other) == "not_compared"
    assert proof_group(_proved(critic=None, paired_by="code")) == "clouded"
    # R1: the brand prototype is reference only; a difference from it alone is never red.
    assert proof_of(_proved(source="brand_prototype", critic=None))["not_compared"] == "prototype"
    assert proof_group(_proved(source="gv_review", critic=None)) == "not_compared"
    # GV's rule, GV's limits and the vendor's own sum compare nothing drawn by someone else.
    rule = _proved(source="rules", witnesses=["rule"], rule_check_id="C1", critic=None)
    assert not needs_pairing(rule) and proof_group(rule) == "clouded"
    summed = _proved(
        source="vendor_arithmetic",
        witnesses=["arithmetic"],
        parts_sum="80.125",
        inches={"architect": None},
        critic=None,
        created_by="consistency",  # code verified the string's geometry (#1639)
    )
    assert proof_group(summed) == "clouded"
    # A string an AI grouped is never red on its sum alone: only code's geometry proves it.
    assert proof_group({**summed, "created_by": "ledger"}) == "not_compared"
    assert proof_group({**summed, "created_by": "ledger", "string_geometry": "verified"}) == (
        "clouded"
    )
    limits = _proved(
        source="architect",
        witnesses=["rule"],
        inches={"architect": None},
        critic=None,
        settled={"method": "gv_limits", "how": 'GV\'s filler limits are 1" to 2" (x); y.'},
    )
    assert not needs_pairing(limits) and proof_group(limits) == "clouded"
    # The reviewer's Agree is still an override.
    assert proof_group({**unpaired, "decision": {"decision": "agree", "by": "R"}}) == "clouded"


def test_counts_split_the_screen_counts_and_follow_decisions() -> None:
    findings = [_proved(), _ai_only(), _decision(), _right(), _right(compared=False)]
    assert proof_counts(findings) == {
        "clouded": 1,
        "right": 1,
        "rule_note": 0,
        "likely_problem": 0,
        "not_compared": 3,
        "dismissed": 0,
    }
    decided = [_proved(decision={"decision": "not_a_problem", "by": "R"}), _ai_only()]
    assert proof_counts(decided)["dismissed"] == 1 and proof_counts(decided)["clouded"] == 0


# ---- colours and words ---------------------------------------------------------------------------


def test_red_and_green_read_on_white_and_stay_dark_in_greyscale() -> None:
    assert contrast(RED) >= 4.5
    assert contrast(GREEN) >= 4.5
    for colour in (RED, GREEN):
        grey = 0.299 * colour.red + 0.587 * colour.green + 0.114 * colour.blue
        assert grey < 0.45  # prints as a dark grey, never washed out
    assert contrast(Color(1, 1, 1)) == pytest.approx(1.0)


def test_the_words_say_what_is_wrong_and_what_it_should_be() -> None:
    finding = _proved(vs_page=2, arch_box={**_box(ARCHITECT_37), "page": 2})
    assert vendor_value(finding) == '46"'
    assert should_be(finding) == '37" (architect 3\'-1", page 2)'
    assert proof_line(finding) == (
        "A blind second pick chose the same architect dimension and code checked both ends; "
        "printed at both spots, read by code; a second reading agrees; checked on the drawing"
    )
    summed = _proved(
        source="vendor_arithmetic",
        parts=['40"', '40 1/8"'],
        printed_total='80"',
        parts_sum="80.125",
        witnesses=["arithmetic"],
        inches={"architect": None, "vendor": None},
    )
    assert vendor_value(summed) == '80"'
    assert should_be(summed) == "80 1/8\" (sum of the vendor's own parts)"
    assert proof_line(summed).startswith('Parts add up to 80 1/8", printed 80"')
    filler = _proved(
        source="rules",
        architect="GV rule 1-2",
        vendor='5" filler each end',
        witnesses=["rule"],
        inches={"architect": None, "vendor": None},
        settled={
            "outcome": "wrong",
            "method": "gv_limits",
            "how": 'GV\'s filler limits are 1" to 2" (rules.md, the settings in force for '
            'this set); the vendor draws 5" fillers.',
            "was": "check",
        },
    )
    assert vendor_value(filler) == '5"'
    assert should_be(filler) == '1" to 2" (GV filler rule)'
    assert governs(filler) == "GV filler rule"
    assert proof_line(filler).startswith("Settled by GV Review: GV's filler limits")
    nothing = proof_of(_decision())
    assert nothing["group"] == "not_compared" and nothing["rule"] == "R7"
    assert nothing["not_compared"] == "nothing_to_compare"
    assert nothing["reason"] == "architect gives no size or information for it"
    assert nothing["settle"] == "the architect's size or note for the counter overall width"
    assert nothing["should_be"] is None
    assert proof_of(_right())["should_be"] is None
    assert stamp_line("2026-10-18") == f"{STAMP_WORDS}: 2026-10-18"
    assert stamp_line("2026-10-18", ["Ann", "Ann", "Bo"]).endswith("and reviewed by Ann, Bo")


# ---- the rendered drawing ------------------------------------------------------------------------


def _placed(findings: list[dict[str, Any]]) -> list[tuple[int, Any, int | None, Any, str]]:
    out: list[tuple[int, Any, int | None, Any, str]] = []
    for number, finding in enumerate(findings, start=1):
        box = finding.get("vendor_box")
        if box is None:
            out.append((number, finding, None, None, "Not marked on the drawing: no box."))
            continue
        coords = tuple(float(box[k]) for k in ("x0", "y0", "x1", "y1"))
        out.append((number, finding, int(box["page"]), coords, f"Marked on page {box['page']}."))
    return out


def _render(pdf: Path, index: int) -> np.ndarray:
    document = pdfium.PdfDocument(str(pdf))
    try:
        return np.asarray(document[index].render(scale=1).to_pil().convert("RGB")).astype(int)
    finally:
        document.close()


def _count(picture: np.ndarray, colour: Color, box: tuple[float, ...] | None = None) -> int:
    area = picture
    if box is not None:
        x0, y0, x1, y1 = (int(v) for v in box)
        area = picture[max(0, y0) : y1, max(0, x0) : x1]
    target = np.array([colour.red, colour.green, colour.blue]) * 255
    return int((np.abs(area - target).max(axis=2) < 40).sum())


def _words(text: str) -> str:
    """The text with its line breaks as single spaces (a note's coloured pieces extract apart)."""
    return " ".join(text.split())


def _draw(tmp_path: Path, findings: list[dict[str, Any]], **kwargs: Any) -> Path:
    sheet = _sheet(tmp_path / "sheet.pdf")
    out = tmp_path / "proof.pdf"
    proof_markup_pdf(
        sheet,
        _placed(findings),
        out,
        set_name="sheet.pdf",
        stamp=stamp_line("2026-10-18", kwargs.pop("reviewers", ())),
        **kwargs,
    )
    return out


def test_the_drawing_clouds_only_what_code_proved_in_red_and_green(tmp_path: Path) -> None:
    findings = [_proved(), _ai_only(), _decision(), _right()]
    out = _draw(tmp_path, findings)
    reader = PdfReader(str(out))
    # Cover, the sheet, one proof page, notes.
    assert list(reader.page_labels) == ["Cover", "1", "Proof 1", "Notes 1"]
    cover = reader.pages[0].extract_text()
    for words in (
        "Checked automatically by GV Review: 2026-10-18",
        "1\nWrong",
        "1\nRight",
        "2\nNot compared",
        RULE_LINE[:60],
        "Buffet outlet height",
    ):
        assert words in cover, words
    assert "draft" not in cover.lower() and "AI markup" not in cover
    sheet = _words(reader.pages[1].extract_text())
    # The note says what is wrong and what it should be, never cut (#1804).
    assert 'Buffet outlet height 46" should be 37"' in sheet and "…" not in sheet
    assert "Base cabinet height" not in sheet and '12" should be' not in sheet
    picture = _render(out, 1)
    x, y = VENDOR_46
    assert _count(picture, RED, (x - 20, y - 25, x + 40, y + 15)) > 20  # the cloud
    x, y = VENDOR_12
    assert _count(picture, RED, (x - 20, y - 25, x + 40, y + 15)) == 0  # AI reading only
    assert _count(picture, GREEN) > 20  # the "Should be" words
    proof = _words(reader.pages[2].extract_text())
    # The proof card (#1804): the C-number and title, one plain sentence, the two snippets'
    # captions, the reference and how it was checked, in the vendor's words.
    for words in (
        "1 Buffet outlet height",
        (
            'Vendor shows 46" for the buffet outlet height; the architect shows 37" '
            '(printed 3\'-1"). Change it to 37" (9" over).'
        ),
        "VENDOR'S DRAWING · PAGE 1",
        "ARCHITECT DRAWING · SAME SHEET",
        'Vendor shows 46" Architect shows 37"',
        "Reference Architect drawing, on the same sheet",
        (
            "Checked Proved by code: both numbers are printed on the drawings and were read from "
            "the drawings' text; a second, separate check matched the same two dimensions and "
            "both ends"
        ),
    ):
        assert words in proof, words
    for gone in ("blind", "advocate", "Proved by:", "AI"):
        assert gone not in proof, gone
    proof_picture = _render(out, 2)
    assert _count(proof_picture, GREEN) > 50 and _count(proof_picture, RED) > 50  # the rings
    notes = " ".join(reader.pages[3].extract_text().split())
    # What could not be compared is not on the drawing: the notes point to the separate list.
    assert "Not compared (2): listed with the reason" in notes
    assert "Base cabinet height" not in notes
    assert "Right (1)" in notes and "Toe kick height" in notes
    words = "".join(page.extract_text() for page in reader.pages)
    for gone in ("check before sending", "Your decision", "Vendor to confirm", "Confirm:"):
        assert gone not in words, gone


def test_decisions_override_and_the_stamp_names_the_reviewer(tmp_path: Path) -> None:
    findings = [
        _proved(decision={"decision": "not_a_problem", "by": "Ann"}),
        _ai_only(decision={"decision": "agree", "by": "Ann"}),
    ]
    out = _draw(tmp_path, findings, reviewers=["Ann"])
    reader = PdfReader(str(out))
    assert "2026-10-18, and reviewed by Ann" in reader.pages[0].extract_text()
    sheet = _words(reader.pages[1].extract_text())
    assert '12" should be 34"' in sheet and '46" should be' not in sheet
    words = "".join(page.extract_text() for page in reader.pages)
    # The reviewer's Agree says it is GV's reviewer's decision, never by login (#1803, #1804).
    assert "Decided by GV's reviewer" in words and "Agreed by Ann" not in words
    # The reviewer's "Not a problem" is Right, noted "(reviewer)": no fourth group (#1766).
    assert "Right (1)" in words and "Right (reviewer)" in words
    assert "Not a problem" not in words


def test_a_drawing_with_nothing_proved_still_has_its_cover_and_lists(tmp_path: Path) -> None:
    out = _draw(tmp_path, [_ai_only(), _decision()], assumptions=["page 2 is the vendor's."])
    reader = PdfReader(str(out))
    assert list(reader.page_labels) == ["Cover", "1", "Notes 1"]
    cover = reader.pages[0].extract_text()
    assert "None: code proved no wrong value on this set." in cover
    assert "Assumed: page 2 is the vendor's." in cover


def test_the_proof_crops_leave_out_earlier_review_markup(tmp_path: Path) -> None:
    """GV's yellow box over the vendor's value is not in the proof crop (it is cut from the copy
    without earlier review markup), and the copy is removed afterwards."""
    sheet = _sheet(tmp_path / "sheet.pdf")
    x, y = VENDOR_46
    with pikepdf.open(sheet, allow_overwriting_input=True) as pdf:
        stream = pdf.make_stream(b"1 1 0 rg 0 0 60 30 re f")
        stream["/Type"] = pikepdf.Name.XObject
        stream["/Subtype"] = pikepdf.Name.Form
        stream["/BBox"] = pikepdf.Array([0, 0, 60, 30])
        square = pdf.make_indirect(
            pikepdf.Dictionary(
                Type=pikepdf.Name.Annot,
                Subtype=pikepdf.Name.Square,
                Rect=pikepdf.Array([x - 20, HEIGHT - y - 15, x + 40, HEIGHT - y + 15]),
                AP=pikepdf.Dictionary(N=stream),
            )
        )
        pdf.pages[0].obj["/Annots"] = pikepdf.Array([square])
        pdf.save(sheet)
    out = tmp_path / "proof.pdf"
    proof_markup_pdf(sheet, _placed([_proved()]), out, set_name="sheet.pdf", stamp="x")
    yellow = Color(1, 1, 0)
    assert _count(_render(out, 2), yellow) == 0
    assert not list(tmp_path.glob("*.clean.pdf"))


# ---- the plan check never waits ----------------------------------------------------------------


def test_unsure_page_map_questions_get_code_s_likely_answer(tmp_path: Path) -> None:
    from app.review.assistant.full_check_plan import plan_assumptions, read_page_map

    page_map = {
        "files": ["set.pdf"],
        "pages": [
            {
                "file": "set.pdf",
                "page": 28,
                "area": "Buffet",
                "halves": [
                    {"role": "architect", "box": [0, 0, 100, 100]},
                    {"role": "general", "box": [0, 200, 100, 300]},
                ],
            }
        ],
    }
    (tmp_path / "page_map.json").write_text(json.dumps(page_map), encoding="utf-8")
    assumed = plan_assumptions(tmp_path, "set.pdf")
    assert assumed == ["page 28, the lower drawing (no heading) is the vendor's drawing."]
    plain = read_page_map(tmp_path)
    run = read_page_map(tmp_path, assume=True, main_file="set.pdf")
    assert plain is not None and run is not None
    assert plain["pages"][0]["halves"][1]["role"] == "general"  # the popover still asks
    assert run["pages"][0]["halves"][1] == {
        "role": "vendor",
        "box": [0, 200, 100, 300],
        "heading_by": "code",
    }
    # A reviewer's answer wins over code's assumption, and is no longer assumed.
    (tmp_path / "plan_answers.jsonl").write_text(
        json.dumps({"id": "role:f0:p28:h1", "answer": "general"}) + "\n", encoding="utf-8"
    )
    assert plan_assumptions(tmp_path, "set.pdf") == []
    answered = read_page_map(tmp_path, assume=True, main_file="set.pdf")
    assert answered is not None and answered["pages"][0]["halves"][1]["role"] == "general"


def test_the_vendor_s_own_string_or_span_is_not_labelled_architect(tmp_path: Path) -> None:
    """A proof that is the vendor's own string (parts against total) or the vendor printing the
    same span twice has no architect's side: the proof page says whose numbers they are."""
    summed = _proved(
        id="F5",
        item="Run height string",
        source="vendor_arithmetic",
        architect='80"',
        vendor='80"',
        parts=['40"', '40 1/8"'],
        printed_total='80"',
        parts_sum="80.125",
        witnesses=["arithmetic"],
        inches={"architect": None, "vendor": "80"},
        arch_box=None,
        critic=None,
        created_by="consistency",
    )
    span = _proved(
        id="F6",
        item="Base: same span printed twice",
        source="cross_view",
        reference_side="vendor",
        created_by="consistency",
        architect='34"',
        vendor='46"',
        inches={"architect": "34", "vendor": "46"},
        critic=None,
    )
    assert vendor_own(summed) == "sum" and vendor_own(span) == "span"
    assert vendor_own(_proved()) is None
    # The same span proved by the drawing's geometry needs no other pairing proof.
    assert not needs_pairing(span) and proof_group(span) == "clouded"
    assert needs_pairing({**span, "source": "architect", "rule_check_id": None})
    out = _draw(tmp_path, [summed, span])
    proof = _words("".join(page.extract_text() for page in PdfReader(str(out)).pages))
    # The sum written out (#1804), and the right panel says whose numbers they are.
    assert "THE VENDOR'S OWN PARTS, ADDED UP" in proof
    assert 'Sum 40" + 40 1/8" = 80 1/8", not the 80" printed' in proof
    assert (
        'Vendor prints 80" for the run height string, but its own parts add up to 80 1/8"' in proof
    )
    assert "VENDOR'S OTHER DIMENSION OF THE SAME SPAN" in proof
    assert "Reference The vendor's other dimension of the same span" in proof
    assert "ARCHITECT DRAWING" not in proof


def test_two_vendor_pages_that_disagree_name_no_should_be(tmp_path: Path) -> None:
    """The same item in two sizes on two of the vendor's drawings (#1906): the card shows both,
    and says one of them is wrong without picking which."""
    pages = _proved(
        id="F7",
        item="Base B-30 width",
        source="cross_view",
        reference_side="vendor",
        created_by="ledger",
        architect='37"',
        vendor='46"',
        inches={"architect": "37", "vendor": "46"},
        critic=None,
        same_item_pages={"pages": [{"page": 1, "value": '37"'}, {"page": 1, "value": '46"'}]},
    )
    assert vendor_own(pages) == "pages" and proof_group(pages) == "clouded"
    out = _draw(tmp_path, [pages])
    proof = _words("".join(page.extract_text() for page in PdfReader(str(out)).pages))
    assert "VENDOR'S OTHER DRAWING OF THE SAME ITEM" in proof
    assert "correct one so both pages agree" in proof
    assert "Change it to" not in proof and "ARCHITECT DRAWING" not in proof


# ---- the Not-compared list (#1619) -----------------------------------------------------------------


def test_the_not_compared_list_has_the_same_rows_as_the_workbook_and_the_screen(
    tmp_path: Path,
) -> None:
    import io as stdio

    import openpyxl

    from app.review.assistant import not_compared
    from app.review.assistant.proof_drawing import not_compared_pdf

    findings = [_proved(), _ai_only(), _decision(), _right(), _right(id="F5", compared=False)]
    rows = not_compared.rows(findings, areas={1: "Lobby"})
    # The screen counts the same rows: every Not-compared item (the proof group).
    assert len(rows) == proof_counts(findings)["not_compared"] == 3
    assert [r.reason for r in rows] == ["nothing_to_compare", "unreadable", "not_seen"]
    sheet = _sheet(tmp_path / "sheet.pdf")
    out = tmp_path / "list.pdf"
    listed = not_compared_pdf(sheet, rows, out, set_name="sheet.pdf", stamp="Checked")
    assert listed == len(rows)
    words = " ".join("".join(p.extract_text() for p in PdfReader(str(out)).pages).split())
    assert "3 items could not be compared" in words
    # The trade's words next to ours (industry review, topic 5).
    assert "Not reviewed: not shown on the architect's drawings (1)" in words
    assert "architect gives no size or information" in words and "Illegible (1)" in words
    assert "Counter overall width" in words and "the architect's size or note for" in words
    book = openpyxl.load_workbook(stdio.BytesIO(not_compared.xlsx(rows)))
    assert book.sheetnames == ["Not compared", "Found across the set", "Summary"]
    table = [[c.value for c in row] for row in book["Not compared"].iter_rows()]
    assert table[0] == list(not_compared.COLUMNS)
    assert [line[4] for line in table[1:]] == [r.item for r in rows]
    assert table[1][2] == "Lobby" and table[1][6] == "no width printed"
    assert book["Not compared"]["A2"].number_format == "@"  # text, so 1/8 stays 1/8
    summary = [[c.value for c in row] for row in book["Summary"].iter_rows()]
    assert ["Total", "items of this check", "3"] in summary
    # The set's own gaps code found (#1623 finish check, #1584) are listed apart, never counted as
    # items: the screen's count stays the item rows'.
    gaps = not_compared.gap_rows(
        [
            {
                "page": 4,
                "what": "Lobby: architect's QZ-1 against the vendor's SL-2",
                "why": "not_comparable",
                "needs": "a printed legend",
            },
            {
                "page": 9,
                "what": "Lobby: round 2 does not include this area",
                "why": "missing_round",
                "needs": "the round-2 sheet",
            },
        ]
    )
    assert [(g.reason, g.level) for g in gaps] == [
        ("finish_not_comparable", "set"),
        ("no_sheet", "set"),
    ]
    out2 = tmp_path / "list2.pdf"
    assert not_compared_pdf(sheet, [*rows, *gaps], out2, set_name="sheet.pdf", stamp="Checked") == 3
    words2 = " ".join("".join(p.extract_text() for p in PdfReader(str(out2)).pages).split())
    assert "Found by code across the set (2)" in words2
    assert "Not reviewed: finish codes not comparable (1)" in words2
    assert not_compared.csv_text(rows).count("\r\n") == 4
    # The workbook's first sheet holds the items only (the screen's count); the set's gaps are a
    # sheet of their own (the app test: 46 on screen against 1,232 rows in one sheet).
    both = openpyxl.load_workbook(stdio.BytesIO(not_compared.xlsx([*rows, *gaps])))
    assert both["Not compared"].max_row == 1 + 3
    assert both["Found across the set"].max_row == 1 + 2
    # The product scope: an item of a product the check was not run for is left out.
    scoped = not_compared.rows(
        [_decision(category="cabinet"), _decision(id="F9", category="countertop")],
        products=["countertop"],
        known_products=["countertop", "cabinet", "other"],
    )
    assert [r.finding_id for r in scoped] == ["F9"]
