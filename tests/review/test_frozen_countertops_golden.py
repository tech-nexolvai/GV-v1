"""A countertop section frozen at sign-off must stay printable by every later version (#1182).

`tests/fixtures/signed_snapshots/countertops_v1.json` is a frozen section exactly as an approval
stores it today (synthetic data only). It is committed and never regenerated: when a change to the
screen's `CountertopResultsOut` makes it unreadable, this test fails, and the change must either keep
the old shape readable (a default for a new field, an old `Literal` word kept) or raise
`COUNTERTOP_SECTION_VERSION` with an upgrade step. Otherwise every review signed before the change
could no longer be published.
"""

from __future__ import annotations

import io
import json
from pathlib import Path
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from app.review.signed_exports import (
    COUNTERTOP_SECTION_VERSION,
    FrozenCountertops,
    SignedExportRefused,
    printable_countertops,
)
from reports.findings_pdf import FindingsPdfInput, write_findings_pdf
from reports.spreadsheet import StoredFinding, write_stored_workbook

FIXTURE = Path(__file__).parents[1] / "fixtures" / "signed_snapshots" / "countertops_v1.json"


FINDING = StoredFinding(
    rule_id="CT-WIDTH-001",
    outcome="REVIEW_REQUIRED",
    severity="FLAG",
    snapshot_id="synthetic-snapshot",
    engine_version="0.1.0",
    trace={},
    reason="Synthetic: held for the reviewer.",
)


def _frozen() -> FrozenCountertops:
    return FrozenCountertops.model_validate(json.loads(FIXTURE.read_text()))


def test_the_committed_v1_section_still_loads() -> None:
    printable = printable_countertops(_frozen())

    assert _frozen().schema_version == 1 <= COUNTERTOP_SECTION_VERSION
    assert [item.page_number for item in printable.items] == [1, 2]
    assert printable.items[0].reviewer_decision is not None
    assert [page.page_number for page in printable.pages_without_countertop] == [3]
    assert [row.page_number for row in printable.rows_not_checked] == [2]


def test_the_committed_v1_section_still_prints_in_the_pdf_and_the_workbook() -> None:
    printable = printable_countertops(_frozen())

    pdf = write_findings_pdf(
        FindingsPdfInput(
            package_revision_id=uuid4(),
            revision_number=1,
            vendor="Synthetic vendor",
            findings=(FINDING,),
            countertop_results=printable.items,
            pages_without_countertop=printable.pages_without_countertop,
            rows_not_checked=printable.rows_not_checked,
        )
    )
    workbook = write_stored_workbook(
        (FINDING,),
        countertop_results=printable.items,
        pages_without_countertop=printable.pages_without_countertop,
        rows_not_checked=printable.rows_not_checked,
    )

    text = " ".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(pdf)).pages)
    assert "Countertop row on page 1" in text
    assert "Synthetic: both readers found no countertop line" in text
    sheet = load_workbook(io.BytesIO(workbook))["Countertops"]
    labels = {cell.value for line in sheet.iter_rows() for cell in line}
    assert {"Countertop row on page 1", "Countertop row on page 2"} <= labels


def test_a_section_the_code_cannot_read_is_refused_with_a_reason_not_a_crash() -> None:
    broken = FrozenCountertops(schema_version=1, results={"items": "not a list"})
    unknown = FrozenCountertops(schema_version=COUNTERTOP_SECTION_VERSION + 1, results={})

    with pytest.raises(SignedExportRefused, match="cannot be read"):
        printable_countertops(broken)
    with pytest.raises(SignedExportRefused, match="does not know"):
        printable_countertops(unknown)
