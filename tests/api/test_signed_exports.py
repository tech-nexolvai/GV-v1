"""The handed-over record is the approval's record, never a stale or partial report."""

import io
from datetime import UTC, datetime, timedelta
from typing import Any

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import CheckRun, Finding, OutputArtifact, Page
from app.models.review import Approval, ApprovedFinding, ReviewAction, ReviewActionKind
from app.models.signed_exports import ApprovalExportBundle, ApprovalExportSnapshot
from app.review.session import open_session, record_action
from app.review.signed_exports import canonical, deterministic_facts, request_signed_exports
from storage.local import LocalStore
from tests.api.test_v1_loop import PROJECT, _uploaded
from tests.api.test_v1_loop import client as client_fixture
from tests.api.test_v1_loop import session as session_fixture
from tests.api.test_v1_loop import store as store_fixture
from tests.review.test_session import _finding
from workflow.signed_outputs import generate_signed_outputs

session = session_fixture
store = store_fixture
client = client_fixture


def _legacy(db: Session, storage: LocalStore) -> tuple[Any, Approval, list[Finding]]:
    package_id, revision, _ = _uploaded(db, storage)
    revision.state = "APPROVED"
    from app.models import DocumentVersion

    version = db.scalars(select(DocumentVersion)).one()
    db.add(
        Page(
            document_version_id=version.id,
            index=0,
            content_hash="a" * 64,
            width_pt=200,
            height_pt=100,
            rotation=0,
            has_vector_text=True,
            media_box=["0", "0", "200", "100"],
            crop_box=["0", "0", "200", "100"],
        )
    )
    failed = _finding(db, revision)
    findings = [failed]
    run = db.get(CheckRun, failed.check_run_id)
    assert run is not None
    for _ in range(8):
        other = CheckRun(
            package_revision_id=revision.id,
            rule_snapshot_id=run.rule_snapshot_id,
            engine_version=run.engine_version,
        )
        db.add(other)
        db.flush()
        finding = Finding(
            check_run_id=other.id,
            package_revision_id=revision.id,
            outcome="NOT_FOUND",
            severity="FLAG",
            trace={},
            parameter_set_versions={},
            reason="Synthetic source information not supplied",
        )
        db.add(finding)
        db.flush()
        findings.append(finding)
    review = open_session(db, package_revision_id=revision.id, reviewer="ana@example.com")
    for finding in findings:
        record_action(
            db,
            review_session_id=review.id,
            finding_id=finding.id,
            action=(
                ReviewActionKind.CONFIRM if finding.outcome == "FAIL" else ReviewActionKind.DISMISS
            ),
            actor="ana@example.com",
            note="Not checkable: synthetic source missing",
        )
    approval = Approval(package_revision_id=revision.id, approved_by="ana@example.com")
    db.add(approval)
    db.flush()
    db.add_all(
        ApprovedFinding(approval_id=approval.id, package_revision_id=revision.id, finding_id=f.id)
        for f in findings
    )
    review.completed_at = approval.created_at
    db.commit()
    return package_id, approval, findings


def _text(content: bytes) -> str:
    return " ".join(
        " ".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(content)).pages).split()
    )


def test_legacy_explicit_request_publishes_all_three_and_never_changes_outcomes(
    session: Session, store: LocalStore, client: Any, monkeypatch: pytest.MonkeyPatch
) -> None:
    package_id, approval, findings = _legacy(session, store)
    base = f"/api/v1/projects/{PROJECT}/packages/{package_id}"
    before = canonical(deterministic_facts(session, approval))
    assert client.get(base + "/signed-exports").json()["status"] == "not_requested"
    assert session.scalar(select(ApprovalExportSnapshot)) is None
    # Old artifacts exist and look plausible, but are not an approved publication.
    for kind in ("findings_pdf", "findings_workbook", "redline"):
        saved = store.put(
            "before/" + kind, io.BytesIO(b"before review"), content_type="application/pdf"
        )
        session.add(
            OutputArtifact(
                package_revision_id=approval.package_revision_id,
                kind=kind,
                storage_key=saved.key,
                sha256=saved.sha256,
                media_type="application/pdf",
                findings=9,
            )
        )
    session.commit()
    for endpoint in ("report", "report.pdf", "redline.pdf"):
        assert client.get(base + "/" + endpoint).status_code == 409
    requested = client.post(base + "/signed-exports")
    assert requested.status_code == 202, requested.text
    assert requested.json()["status"] == "preparing"
    from workflow.changed_values import ChangedValues

    monkeypatch.setattr(
        "workflow.changed_values.changed_values_for_revision",
        lambda *args, **kwargs: ChangedValues(
            "available", None, ("LATER SETTING MUST NOT REPLACE PINNED REPORT",), ()
        ),
    )
    for endpoint in ("report", "report.pdf", "redline.pdf"):
        assert client.get(base + "/" + endpoint).status_code == 409
    bundle = generate_signed_outputs(session, store, approval.id)
    session.commit()
    assert generate_signed_outputs(session, store, approval.id).id == bundle.id
    session.commit()
    assert len(session.scalars(select(ApprovalExportBundle)).all()) == 1
    assert client.get(base + "/signed-exports").json()["status"] == "ready"
    for endpoint in ("report.pdf", "redline.pdf"):
        response = client.get(base + "/" + endpoint)
        assert response.status_code == 200, response.text
        text = _text(response.content)
        assert (
            text.count("Not checked: dismissed as not checkable by reviewer ana@example.com") == 8
        )
        assert "FAIL: problem for the vendor, confirmed by reviewer ana@example.com" in text
        assert "Reason: Not checkable: synthetic source missing" in text
        assert str(approval.id) in text and "9 findings covered" in text
        assert "Signed off by ana@example.com" in text
        assert "outcome (unchanged): PASS" not in text
        assert "LATER SETTING MUST NOT REPLACE PINNED REPORT" not in text
    response = client.get(base + "/report")
    assert response.status_code == 200
    book = load_workbook(io.BytesIO(response.content))
    review_text = " ".join(str(c.value or "") for row in book["Signed review"] for c in row)
    assert (
        review_text.count("Not checked: dismissed as not checkable by reviewer ana@example.com")
        == 8
    )
    assert "FAIL: problem for the vendor, confirmed by reviewer ana@example.com" in review_text
    assert str(approval.id) in review_text and "9 findings covered" in review_text
    assert "PASS" not in review_text
    assert "LATER SETTING MUST NOT REPLACE PINNED REPORT" not in str(
        [list(sheet.values) for sheet in book]
    )
    assert canonical(deterministic_facts(session, approval)) == before
    assert [f.outcome for f in findings] == ["FAIL"] + ["NOT_FOUND"] * 8
    # A later run cannot change the earlier publication or its pinned contents.
    earlier_pdf = client.get(base + "/report.pdf").content
    run = session.get(CheckRun, findings[0].check_run_id)
    assert run is not None
    run.superseded_at = datetime.now(UTC)
    session.commit()
    assert client.get(base + "/report.pdf").content == earlier_pdf


def test_legacy_ambiguous_history_refuses_plainly(
    session: Session, store: LocalStore, client: Any
) -> None:
    package_id, _approval, findings = _legacy(session, store)
    action = session.scalars(
        select(ReviewAction).where(ReviewAction.finding_id == findings[0].id)
    ).one()
    session.add(
        ReviewAction(
            review_session_id=action.review_session_id,
            package_revision_id=action.package_revision_id,
            finding_id=action.finding_id,
            action="dismiss",
            actor=action.actor,
            note="Second action at an indistinguishable time",
            created_at=action.created_at,
        )
    )
    session.commit()
    response = client.post(f"/api/v1/projects/{PROJECT}/packages/{package_id}/signed-exports")
    assert response.status_code == 409
    assert "ambiguous" in response.text
    assert session.scalar(select(ApprovalExportSnapshot)) is None


def test_render_failure_never_publishes_partial_files(
    session: Session, store: LocalStore, client: Any, monkeypatch: pytest.MonkeyPatch
) -> None:
    package_id, approval, _ = _legacy(session, store)
    request_signed_exports(session, approval.id)
    session.commit()

    def broken(*args: Any, **kwargs: Any) -> bytes:
        raise OSError("synthetic workbook render failure")

    monkeypatch.setattr("workflow.signed_outputs.write_stored_workbook", broken)
    with pytest.raises(OSError, match="synthetic"):
        generate_signed_outputs(session, store, approval.id)
    session.rollback()
    assert session.scalar(select(ApprovalExportBundle)) is None
    from sqlalchemy.orm import sessionmaker

    from workflow.signed_outputs import record_publication_failure

    record_publication_failure(
        sessionmaker(bind=session.get_bind()), approval.id, OSError("private source text")
    )
    response = client.get(f"/api/v1/projects/{PROJECT}/packages/{package_id}/signed-exports")
    assert response.json()["status"] == "failed"
    assert "private source text" not in response.text
    for endpoint in ("report", "report.pdf", "redline.pdf"):
        assert (
            client.get(f"/api/v1/projects/{PROJECT}/packages/{package_id}/{endpoint}").status_code
            == 409
        )


def test_later_approval_cannot_download_an_earlier_bundle(
    session: Session, store: LocalStore, client: Any
) -> None:
    package_id, approval, findings = _legacy(session, store)
    request_signed_exports(session, approval.id)
    generate_signed_outputs(session, store, approval.id)
    later = Approval(
        package_revision_id=approval.package_revision_id,
        approved_by="second reviewer",
        created_at=approval.created_at + timedelta(seconds=1),
    )
    session.add(later)
    session.flush()
    session.add_all(
        ApprovedFinding(
            approval_id=later.id, finding_id=f.id, package_revision_id=approval.package_revision_id
        )
        for f in findings
    )
    session.commit()
    response = client.get(f"/api/v1/projects/{PROJECT}/packages/{package_id}/report.pdf")
    assert response.status_code == 409
    assert "not been requested" in response.text


def _outbox_requests(db: Session, approval_id: Any) -> int:
    from app.models.outbox import OutboxEntry
    from app.review.signed_exports import WORKFLOW

    return sum(
        1
        for entry in db.scalars(select(OutboxEntry).where(OutboxEntry.workflow == WORKFLOW))
        if entry.payload.get("approval_id") == str(approval_id)
    )


def test_ready_status_lists_each_signed_file_with_its_recorded_size(
    session: Session, store: LocalStore, client: Any, monkeypatch: pytest.MonkeyPatch
) -> None:
    package_id, approval, _ = _legacy(session, store)
    base = f"/api/v1/projects/{PROJECT}/packages/{package_id}"
    assert client.post(base + "/signed-exports").json()["files"] is None
    assert client.get(base + "/signed-exports").json()["files"] is None
    bundle = generate_signed_outputs(session, store, approval.id)
    session.commit()

    # The status call answers from the database: it never opens a stored file.
    def no_reading(*args: Any, **kwargs: Any) -> Any:
        raise AssertionError("the status call must not read stored files")

    monkeypatch.setattr(LocalStore, "get", no_reading)
    body = client.get(base + "/signed-exports").json()
    monkeypatch.undo()
    assert body["status"] == "ready"
    stored = {
        kind: session.get(OutputArtifact, artifact_id)
        for kind, artifact_id in (
            ("findings_pdf", bundle.pdf_id),
            ("workbook", bundle.workbook_id),
            ("redline", bundle.redline_id),
        )
    }
    assert [item["kind"] for item in body["files"]] == ["findings_pdf", "workbook", "redline"]
    for item in body["files"]:
        artifact = stored[item["kind"]]
        assert artifact is not None
        assert item["bytes"] == len(store.get(artifact.storage_key).read()) > 0
        assert item["media_type"] == artifact.media_type
    assert client.post(base + "/signed-exports").json()["files"] == body["files"]


def test_an_older_bundle_without_a_recorded_size_says_null(
    session: Session, store: LocalStore, client: Any
) -> None:
    package_id, approval, findings = _legacy(session, store)
    snapshot = request_signed_exports(session, approval.id)
    ids = {}
    for kind in ("findings_pdf", "findings_workbook", "redline"):
        saved = store.put(f"older/{kind}", io.BytesIO(b"older bundle"), content_type="x/y")
        artifact = OutputArtifact(
            package_revision_id=approval.package_revision_id,
            kind=kind,
            storage_key=saved.key,
            sha256=saved.sha256,
            media_type="application/pdf",
            findings=len(findings),
        )
        session.add(artifact)
        session.flush()
        ids[kind] = artifact.id
    session.add(
        ApprovalExportBundle(
            snapshot_id=snapshot.id,
            package_revision_id=approval.package_revision_id,
            pdf_id=ids["findings_pdf"],
            workbook_id=ids["findings_workbook"],
            redline_id=ids["redline"],
        )
    )
    session.commit()
    body = client.get(f"/api/v1/projects/{PROJECT}/packages/{package_id}/signed-exports").json()
    assert body["status"] == "ready"
    assert [item["bytes"] for item in body["files"]] == [None, None, None]


def test_retry_after_a_failure_requeues_once_for_the_same_approval_and_snapshot(
    session: Session, store: LocalStore, client: Any, monkeypatch: pytest.MonkeyPatch
) -> None:
    from sqlalchemy.orm import sessionmaker

    from workflow.signed_outputs import record_publication_failure

    package_id, approval, findings = _legacy(session, store)
    base = f"/api/v1/projects/{PROJECT}/packages/{package_id}/signed-exports"
    first = client.post(base)
    assert first.json()["status"] == "preparing"
    snapshot = session.scalars(select(ApprovalExportSnapshot)).one()
    facts = canonical(deterministic_facts(session, approval))
    assert _outbox_requests(session, approval.id) == 1
    # Asking again while the first request is still being prepared queues nothing.
    assert client.post(base).json()["status"] == "preparing"
    assert _outbox_requests(session, approval.id) == 1

    factory = sessionmaker(bind=session.get_bind())
    record_publication_failure(factory, approval.id, OSError("synthetic"))
    session.expire_all()
    assert client.get(base).json()["status"] == "failed"

    retried = client.post(base)
    assert retried.status_code == 202, retried.text
    assert retried.json() == {
        "approval_id": str(approval.id),
        "status": "preparing",
        "files": None,
    }
    assert _outbox_requests(session, approval.id) == 2
    assert client.get(base).json()["status"] == "preparing"
    # A second press while the retry is preparing does not queue twice.
    assert client.post(base).json()["status"] == "preparing"
    assert _outbox_requests(session, approval.id) == 2

    # The retry failed too: one more press, one more request.
    record_publication_failure(factory, approval.id, OSError("synthetic"))
    session.expire_all()
    assert client.get(base).json()["status"] == "failed"
    assert client.post(base).json()["status"] == "preparing"
    assert _outbox_requests(session, approval.id) == 3

    # Same approval, same frozen snapshot, nothing re-signed, no outcome changed.
    session.expire_all()
    assert session.scalars(select(ApprovalExportSnapshot)).one().id == snapshot.id
    assert [a.id for a in session.scalars(select(Approval))] == [approval.id]
    assert canonical(deterministic_facts(session, approval)) == facts
    assert [session.get(Finding, f.id).outcome for f in findings] == [  # type: ignore[union-attr]
        "FAIL"
    ] + ["NOT_FOUND"] * 8

    generate_signed_outputs(session, store, approval.id)
    session.commit()
    assert client.get(base).json()["status"] == "ready"
    # Once ready, asking again queues nothing.
    assert client.post(base).json()["status"] == "ready"
    assert _outbox_requests(session, approval.id) == 3


def test_retry_is_project_scoped(session: Session, store: LocalStore, client: Any) -> None:
    from uuid import uuid4

    package_id, _approval, _ = _legacy(session, store)
    response = client.post(f"/api/v1/projects/{uuid4()}/packages/{package_id}/signed-exports")
    assert response.status_code == 404


def test_signed_report_lists_a_split_page_and_the_pages_with_no_countertop(
    session: Session, store: LocalStore, client: Any, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The signed PDF and workbook carry what the countertop results say (#1093): the page the AIs
    picked different lines on, held, and the pages both AIs found no countertop line on."""
    from app.schemas.visual_ui import CountertopResultsOut, PageWithoutCountertopOut
    from tests.reports.test_countertop_summary_reports import NONE_REASON, SPLIT_REASON, _split_item

    package_id, approval, _ = _legacy(session, store)
    base = f"/api/v1/projects/{PROJECT}/packages/{package_id}"

    def results(_db: Session, package: Any, revision: Any) -> CountertopResultsOut:
        return CountertopResultsOut(
            package_id=package,
            revision_id=revision.id,
            items=(_split_item(),),
            pages_without_countertop=(PageWithoutCountertopOut(page_number=4, reason=NONE_REASON),),
        )

    monkeypatch.setattr("workflow.signed_outputs._countertop_results_for_revision", results)
    assert client.post(base + "/signed-exports").status_code == 202
    generate_signed_outputs(session, store, approval.id)
    session.commit()

    text = _text(client.get(base + "/report.pdf").content)
    assert "Countertop row on page 3" in text and SPLIT_REASON in text
    assert "PAGES WITH NO COUNTERTOP FOUND" in text and f"Page 4: {NONE_REASON}" in text
    book = load_workbook(io.BytesIO(client.get(base + "/report").content))
    assert SPLIT_REASON in [cell.value for row in book["Countertops"] for cell in row]
    assert [cell.value for cell in book["No Countertop Found"][2]] == ["4", NONE_REASON]


def test_signed_report_lists_a_second_countertop_row_that_was_not_checked(
    session: Session, store: LocalStore, client: Any, monkeypatch: pytest.MonkeyPatch
) -> None:
    """#1108: a second countertop row an AI named on a read page reaches the signed PDF and
    workbook as not checked."""
    from app.schemas.visual_ui import CountertopResultsOut, RowNotCheckedOut
    from tests.reports.test_countertop_summary_reports import SECOND_ROW_REASON

    package_id, approval, _ = _legacy(session, store)
    base = f"/api/v1/projects/{PROJECT}/packages/{package_id}"

    def results(_db: Session, package: Any, revision: Any) -> CountertopResultsOut:
        return CountertopResultsOut(
            package_id=package,
            revision_id=revision.id,
            items=(),
            rows_not_checked=(RowNotCheckedOut(page_number=2, reason=SECOND_ROW_REASON),),
        )

    monkeypatch.setattr("workflow.signed_outputs._countertop_results_for_revision", results)
    assert client.post(base + "/signed-exports").status_code == 202
    generate_signed_outputs(session, store, approval.id)
    session.commit()

    text = _text(client.get(base + "/report.pdf").content)
    assert "SECOND COUNTERTOP ROWS NOT CHECKED" in text
    assert "Page 2: An AI found a second countertop on this page" in text
    book = load_workbook(io.BytesIO(client.get(base + "/report").content))
    assert [cell.value for cell in book["Second Rows Not Checked"][2]] == ["2", SECOND_ROW_REASON]
