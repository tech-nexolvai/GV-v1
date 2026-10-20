"""The reviewer's overrides on a full check's result (#1803), through the app on a real PostgreSQL.

On a Not-compared item the screen offers "It's a mistake: add a red mark" (an Agree with the value
it should be) and "It's fine: mark Right" (Not a problem); on a Wrong item "Not a mistake: remove
the red mark" (Not a problem). Pinned here:

- the reviewer's "should be" is stored with the decision and comes back on the finding;
- it makes a Not-compared item Wrong even where the architect prints no value, with the green
  "should be" marked "(reviewer)" and the proof saying it is the reviewer's decision, never the
  reviewer's login;
- the remade files (cloud drawing, Not-compared PDF and workbook, package) follow the decisions:
  a decided item leaves the Not-compared list, the new red is clouded with its proof entry.

No paid call is possible: the provider is scripted; the sandbox is a fake.
"""

from __future__ import annotations

import io
import zipfile
from typing import Any
from uuid import UUID

import pytest
from pypdf import PdfReader

from app.review.assistant import chat_engine, chat_trust
from tests.api.test_assistant_chats import _events, _named, _world
from tests.api.test_assistant_chats_findings import _red_near
from tests.review.assistant.chat_fakes import ScriptedProvider, reply, text, tool
from tests.review.assistant.test_chat_findings import VENDOR_BOX, finding_input, make_pdf

pytest_plugins = ("tests.app.postgres_fixture",)

PDF3 = make_pdf(3)
PAGE3_BOX = {**VENDOR_BOX, "page": 3}
ASK = {"model": "claude-opus-5-5", "effort": "high", "thinking": True}


@pytest.fixture(autouse=True)
def _no_checks_modules(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(chat_trust, "load_module", lambda name: None)
    monkeypatch.setattr(chat_engine, "load_module", lambda name: None)


def _provider() -> ScriptedProvider:
    """F1: a difference only the AI read (Not compared); F2 on page 3: the architect gives no
    size (Not compared)."""
    return ScriptedProvider(
        lead=[
            reply(
                tool("r1", "report_finding", finding_input()),
                tool(
                    "r2",
                    "report_finding",
                    finding_input(
                        item="Sink base width",
                        status="check",
                        architect="no size printed",
                        vendor='33"',
                        page=3,
                        vendor_box=dict(PAGE3_BOX),
                        why="The architect prints no width for this cabinet",
                    ),
                ),
                stop="tool_use",
            ),
            reply(text("Two items could not be compared.")),
        ]
    )


def _decide(world: Any, chat_id: str, message_id: str, finding_id: str, **body: Any) -> Any:
    return world.client.post(
        f"/api/v1/chats/{chat_id}/messages/{message_id}/findings/{finding_id}/decision", json=body
    )


def _pdf_text(pdf: bytes) -> str:
    return " ".join(
        " ".join((page.extract_text() or "").split()) for page in PdfReader(io.BytesIO(pdf)).pages
    )


def _newest(files: list[dict[str, Any]], ending: str) -> dict[str, Any]:
    return [file for file in files if file["name"].endswith(ending)][-1]


def test_a_not_compared_item_made_wrong_with_the_reviewers_value_and_one_marked_right(postgres_engine, tmp_path) -> None:  # type: ignore[no-untyped-def]
    with _world(postgres_engine, tmp_path, provider=_provider(), pdf=PDF3, pages=3) as world:
        chat_id = world.chat()
        response = world.client.post(
            f"/api/v1/chats/{chat_id}/messages",
            json={**ASK, "mode": "full_check", "products": ["countertop"]},
        )
        assert response.status_code == 200, response.text
        answer = _named(_events(response.text), "message_end")[0]["message"]
        found = {
            b["finding"]["id"]: b["finding"] for b in answer["blocks"] if b["type"] == "finding"
        }
        assert {f["proof"]["group"] for f in found.values()} == {"not_compared"}
        first_list = _pdf_text(
            world.client.get(_newest(answer["files"], "Not compared.pdf")["url"]).content
        )
        assert "Sink base width" in first_list and "Vanity top length" in first_list

        # A "should be" goes with an Agree only; an empty one is no value.
        refused = _decide(
            world, chat_id, answer["id"], "F1", decision="not_a_problem", should_be='5"'
        )
        assert refused.status_code == 422
        # It's a mistake: add a red mark, with the value it should be (the architect gives none).
        made = _decide(
            world,
            chat_id,
            answer["id"],
            "F2",
            decision="agree",
            should_be=' 34 1/2" ',
            note="Hinge side",
        )
        assert made.status_code == 201, made.text
        assert made.json()["should_be"] == '34 1/2"' and made.json()["note"] == "Hinge side"
        # It's fine: mark Right.
        right = _decide(
            world, chat_id, answer["id"], "F1", decision="not_a_problem", note="Field fit"
        )
        assert right.status_code == 201 and right.json().get("should_be") is None

        stored = world.client.get(f"/api/v1/chats/{chat_id}").json()["messages"][-1]
        now = {b["finding"]["id"]: b["finding"] for b in stored["blocks"] if b["type"] == "finding"}
        assert now["F2"]["decision"]["should_be"] == '34 1/2"'
        assert now["F2"]["proof"]["group"] == "clouded"
        assert now["F2"]["proof"]["should_be"] == '34 1/2" (reviewer)'
        assert now["F1"]["proof"]["group"] == "dismissed"

        # The files are made again from the decisions.
        remade = world.client.post(f"/api/v1/chats/{chat_id}/messages/{answer['id']}/cloud-drawing")
        assert remade.status_code == 200, remade.text
        files = world.client.get(f"/api/v1/chats/{chat_id}").json()["messages"][-1]["files"]
        assert len(files) == 2 * len(answer["files"])
        drawing = world.client.get(remade.json()["url"]).content
        words = _pdf_text(drawing)
        assert 'Should be: 34 1/2" (reviewer)' in words
        assert "Decided by GV's reviewer: should be 34 1/2\". Hinge side" in words
        assert "reviewer-one" in words  # the stamp names who reviewed it
        assert "Agreed by" not in words
        assert _red_near(drawing, 4, PAGE3_BOX) > 5  # page 1 is the cover

        # The Not-compared list is made again: both decided items leave it.
        listed = _newest(files, "Not compared.pdf")
        assert listed["id"] not in {f["id"] for f in answer["files"]}
        remade_list = _pdf_text(world.client.get(listed["url"]).content)
        assert "Sink base width" not in remade_list and "Vanity top length" not in remade_list
        sheet = _newest(files, "Not compared.xlsx")
        assert sheet["id"] not in {f["id"] for f in answer["files"]}

        # The package holds the new list and the new red, numbered.
        package = _newest(files, "review package.zip")
        with zipfile.ZipFile(io.BytesIO(world.client.get(package["url"]).content)) as archive:
            inside = archive.namelist()
            letter = _pdf_text(archive.read(next(n for n in inside if "Comment letter" in n)))
            listed_inside = _pdf_text(
                archive.read(next(n for n in inside if n.endswith("Not compared.pdf")))
            )
        assert "C-01" in letter and "Sink base width" in letter and '34 1/2"' in letter
        assert "reviewer-one" not in letter.split("reviewed by")[0]
        assert "Sink base width" not in listed_inside


def _package(world: Any, chat_id: str) -> dict[str, Any]:
    """The answer as the screen reads it, and its newest package's files, each read as the vendor
    would open it: PDFs as text, workbooks as rows."""
    from openpyxl import load_workbook

    message = world.client.get(f"/api/v1/chats/{chat_id}").json()["messages"][-1]
    package = _newest(message["files"], "review package.zip")
    out: dict[str, Any] = {"message": message, "package": package}
    with zipfile.ZipFile(io.BytesIO(world.client.get(package["url"]).content)) as archive:
        for name in archive.namelist():
            data = archive.read(name)
            key = name.rsplit(" - ", 1)[-1]
            if name.endswith(".pdf"):
                out[key] = _pdf_text(data)
                out[key + " bytes"] = data
            elif name.endswith(".xlsx"):
                book = load_workbook(io.BytesIO(data), read_only=True)
                out[key] = {
                    sheet: [
                        " | ".join(str(c or "") for c in row)
                        for row in book[sheet].iter_rows(values_only=True)
                    ]
                    for sheet in book.sheetnames
                }
    return out


def _cells(sheet: dict[str, list[str]]) -> str:
    return " ".join(line for lines in sheet.values() for line in lines)


def test_each_decision_updates_every_file_of_the_package_and_an_undo_restores_it(postgres_engine, tmp_path) -> None:  # type: ignore[no-untyped-def]
    """#1803: decision -> package made again -> download -> each file read: the cloud drawing's new
    red cloud and its proof card, the letter's new comment number, the register, the Not-compared
    PDF and workbook without the item, the cover's counts; the screen is told the old package is
    out of date until it is made again; an undo puts every file back."""
    with _world(postgres_engine, tmp_path, provider=_provider(), pdf=PDF3, pages=3) as world:
        chat_id = world.chat()
        response = world.client.post(
            f"/api/v1/chats/{chat_id}/messages",
            json={**ASK, "mode": "full_check", "products": ["countertop"]},
        )
        assert response.status_code == 200, response.text
        answer = _named(_events(response.text), "message_end")[0]["message"]
        remake = f"/api/v1/chats/{chat_id}/messages/{answer['id']}/cloud-drawing"

        before = _package(world, chat_id)
        assert before["message"]["package_outdated"] is None  # no decision yet
        assert (
            "0 Wrong" in before["GV cloud drawing.pdf"]
            and "2 Not compared" in before["GV cloud drawing.pdf"]
        )
        assert "Sink base width" in before["Not compared.pdf"]
        assert "Sink base width" in _cells(before["Not compared.xlsx"])
        assert (
            "Sink base width" not in _cells(before["Issue register.xlsx"]).split("Not compared")[0]
        )
        assert "C-01" not in before["Comment letter.pdf"]

        # It's a mistake: add a red mark (the architect prints no value: the reviewer gives one).
        made = _decide(
            world,
            chat_id,
            answer["id"],
            "F2",
            decision="agree",
            should_be='34 1/2"',
            note="Hinge side",
        )
        assert made.status_code == 201, made.text
        # The old package is never handed out as if it held the decision.
        assert (
            world.client.get(f"/api/v1/chats/{chat_id}").json()["messages"][-1]["package_outdated"]
            is True
        )
        assert world.client.post(remake).status_code == 200

        after = _package(world, chat_id)
        assert after["message"]["package_outdated"] is False
        assert after["package"]["id"] != before["package"]["id"]
        drawing = after["GV cloud drawing.pdf"]
        assert "1 Wrong" in drawing and "1 Not compared" in drawing  # the cover's counts
        assert "Decided by GV's reviewer: should be 34 1/2\". Hinge side" in drawing  # proof card
        assert "reviewer-one" not in drawing.split("Proof", 1)[1]  # never the login on the card
        assert _red_near(after["GV cloud drawing.pdf bytes"], 4, PAGE3_BOX) > 5  # the new red cloud
        letter = after["Comment letter.pdf"]
        assert "C-01" in letter and "Sink base width" in letter
        assert (
            "GV's reviewer requires 34 1/2\"" in letter
            and "Architect requires 34 1/2" not in letter
        )
        register = _cells(after["Issue register.xlsx"])
        assert "C-01" in register and "Sink base width" in register and '34 1/2"' in register
        assert "Sink base width" not in after["Not compared.pdf"]
        assert "Sink base width" not in _cells(after["Not compared.xlsx"])
        assert "Vanity top length" in after["Not compared.pdf"]  # the undecided item stays listed

        # Undo: the reviewer takes the call back; every file is as code found it again.
        undone = _decide(world, chat_id, answer["id"], "F2", decision="undo")
        assert undone.status_code == 201, undone.text
        refused = _decide(world, chat_id, answer["id"], "F2", decision="undo", should_be='1"')
        assert refused.status_code == 422
        stored = world.client.get(f"/api/v1/chats/{chat_id}").json()["messages"][-1]
        assert stored["package_outdated"] is True
        f2 = next(
            b["finding"]
            for b in stored["blocks"]
            if b["type"] == "finding" and b["finding"]["id"] == "F2"
        )
        assert f2.get("decision") is None and f2["proof"]["group"] == "not_compared"
        assert world.client.post(remake).status_code == 200

        again = _package(world, chat_id)
        assert again["message"]["package_outdated"] is False
        assert (
            "0 Wrong" in again["GV cloud drawing.pdf"]
            and "2 Not compared" in again["GV cloud drawing.pdf"]
        )
        assert "Decided by GV's reviewer" not in again["GV cloud drawing.pdf"]
        assert "C-01" not in again["Comment letter.pdf"]
        assert "Sink base width" in again["Not compared.pdf"]
        assert "Sink base width" in _cells(again["Not compared.xlsx"])
        # The register is back too: no comment for the undone item.
        assert "C-01" not in _cells(again["Issue register.xlsx"])


def _full_check(world: Any) -> tuple[str, dict[str, Any]]:
    chat_id = world.chat()
    response = world.client.post(
        f"/api/v1/chats/{chat_id}/messages",
        json={**ASK, "mode": "full_check", "products": ["countertop"]},
    )
    assert response.status_code == 200, response.text
    return chat_id, _named(_events(response.text), "message_end")[0]["message"]


def test_a_decision_saved_while_the_package_is_made_leaves_it_out_of_date(postgres_engine, tmp_path, monkeypatch) -> None:  # type: ignore[no-untyped-def]
    """#1803 review: the remake read the decisions before a new one landed; its package must
    never look current. The remake also returns every file it made, the package among them."""
    from sqlalchemy.orm import Session

    from app.api import assistant_chats
    from app.review.assistant.chat_decisions import record_decision

    with _world(postgres_engine, tmp_path, provider=_provider(), pdf=PDF3, pages=3) as world:
        chat_id, answer = _full_check(world)
        assert (
            _decide(world, chat_id, answer["id"], "F1", decision="not_a_problem").status_code == 201
        )
        made = assistant_chats.make_decided_files

        def slow(*args: Any, **kwargs: Any) -> Any:
            # Another reviewer's decision lands while the files are being made.
            with Session(postgres_engine) as other:
                record_decision(
                    other,
                    chat_id=UUID(chat_id),
                    message_id=UUID(answer["id"]),
                    finding_id="F2",
                    decision="agree",
                    note=None,
                    decided_by="reviewer-two",
                    should_be='34 1/2"',
                )
                other.commit()
            return made(*args, **kwargs)

        monkeypatch.setattr(assistant_chats, "make_decided_files", slow)
        remade = world.client.post(f"/api/v1/chats/{chat_id}/messages/{answer['id']}/cloud-drawing")
        assert remade.status_code == 200, remade.text
        names = [file["name"] for file in remade.json()["made"]]
        assert remade.json()["id"] == remade.json()["made"][0]["id"]
        assert any(name.endswith("review package.zip") for name in names)
        assert any(name.endswith("Not compared.xlsx") for name in names)
        message = world.client.get(f"/api/v1/chats/{chat_id}").json()["messages"][-1]
        assert message["package_outdated"] is True

        monkeypatch.setattr(assistant_chats, "make_decided_files", made)
        assert (
            world.client.post(
                f"/api/v1/chats/{chat_id}/messages/{answer['id']}/cloud-drawing"
            ).status_code
            == 200
        )
        assert (
            world.client.get(f"/api/v1/chats/{chat_id}").json()["messages"][-1]["package_outdated"]
            is False
        )


def test_a_red_mark_needs_a_number_other_than_the_vendors_own(postgres_engine, tmp_path) -> None:  # type: ignore[no-untyped-def]
    with _world(postgres_engine, tmp_path, provider=_provider(), pdf=PDF3, pages=3) as world:
        chat_id, answer = _full_check(world)
        same = _decide(world, chat_id, answer["id"], "F2", decision="agree", should_be=' 33 " ')
        assert same.status_code == 422 and same.json()["error"] == "should_be_refused"
        words = _decide(world, chat_id, answer["id"], "F2", decision="agree", should_be="wider")
        assert words.status_code == 422 and "number" in words.json()["message"]
        assert (
            _decide(
                world, chat_id, answer["id"], "F2", decision="agree", should_be='34"'
            ).status_code
            == 201
        )


def _older_answer(monkeypatch: pytest.MonkeyPatch) -> None:
    """Answers as stored before #1803: the package block holds no set gaps (stored answers are
    append-only, so the read is changed, not the row)."""
    from app.api import assistant_chats

    read = assistant_chats.package_block

    def older(blocks: Any) -> Any:
        found = read(blocks)
        return None if found is None else {k: v for k, v in found.items() if k != "set_gaps"}

    monkeypatch.setattr(assistant_chats, "package_block", older)


def test_an_older_answers_list_is_made_again_or_refused_never_handed_out_stale(postgres_engine, tmp_path, monkeypatch) -> None:  # type: ignore[no-untyped-def]
    with _world(postgres_engine, tmp_path, provider=_provider(), pdf=PDF3, pages=3) as world:
        chat_id, answer = _full_check(world)
        _older_answer(monkeypatch)
        assert (
            _decide(
                world, chat_id, answer["id"], "F2", decision="agree", should_be='34 1/2"'
            ).status_code
            == 201
        )
        remake = f"/api/v1/chats/{chat_id}/messages/{answer['id']}/cloud-drawing"
        assert world.client.post(remake).status_code == 200
        after = _package(world, chat_id)
        assert "Sink base width" not in after["Not compared.pdf"]
        assert "Sink base width" not in _cells(after["Not compared.xlsx"])

        # A list that cannot be read back is refused, never re-attached as it was.
        sheet = _newest(after["message"]["files"], "Not compared.xlsx")
        folder = next(tmp_path.rglob(sheet["id"]))
        folder.write_bytes(b"not a workbook")
        refused = world.client.post(remake)
        assert refused.status_code == 409 and refused.json()["error"] == "list_not_remade"


def test_an_item_for_gvs_reviewer_only_is_never_decided(postgres_engine, tmp_path, monkeypatch) -> None:  # type: ignore[no-untyped-def]
    from app.api import assistant_chats

    with _world(postgres_engine, tmp_path, provider=_provider(), pdf=PDF3, pages=3) as world:
        chat_id, answer = _full_check(world)
        monkeypatch.setattr(
            assistant_chats.classification, "internal", lambda f: f.get("id") == "F2"
        )
        refused = _decide(world, chat_id, answer["id"], "F2", decision="agree", should_be='34"')
        assert refused.status_code == 422 and refused.json()["error"] == "internal_item"
        assert _decide(world, chat_id, answer["id"], "F2", decision="undo").status_code == 201
        assert (
            _decide(world, chat_id, answer["id"], "F1", decision="not_a_problem").status_code == 201
        )
