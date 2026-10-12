"""The signed files print the countertop results frozen at sign-off, never the live ones (#1182).

The audit (2026-10-11, S1): the signed PDF, workbook and redline rebuilt their countertop section from
the live database at export time, while their findings were frozen at approval. A wall choice saved
after the approval (or before a retried export) reached the signed files, and two pages of one signed
document disagreed. Here a row's wall is changed directly in the database after the approval, the
way a write that slipped past the API would, and the signed files must not notice.
"""

from __future__ import annotations

import io
from collections.abc import Iterator
from pathlib import Path
from typing import Any
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader, PdfWriter
from sqlalchemy import Engine, select
from sqlalchemy.orm import Session

from alembic import command
from app.api.visual_countertops import _countertop_results_for_revision
from app.db.session import session_factory
from app.models import CheckRun, Finding, OutputArtifact, PackageRevision, SlotRowReviewDecision
from app.models.review import Approval, ApprovedFinding
from app.models.signed_exports import ApprovalExportSnapshot
from app.review.signed_exports import (
    load_snapshot,
    printable_countertops,
    request_signed_exports,
)
from storage.local import LocalStore
from tests.api.test_slot_rows import _package_rows, _run_current_checks
from tests.app.postgres_fixture import alembic_config
from workflow.redline_outputs import RedlineOutput
from workflow.signed_outputs import generate_signed_outputs

pytest_plugins = ("tests.app.postgres_fixture",)


@pytest.fixture
def session(postgres_engine: Engine) -> Iterator[Session]:
    config = alembic_config()
    config.attributes["database_url"] = postgres_engine.url.render_as_string(hide_password=False)
    command.upgrade(config, "head")
    opened = session_factory(postgres_engine)()
    try:
        yield opened
    finally:
        opened.close()


def _approved(session: Session, tmp_path: Path) -> tuple[Any, PackageRevision, Approval]:
    """Two checked countertop rows whose walls came from the drawing, signed off."""
    _project, package_id, anchors = _package_rows(
        session, widths_add_up=True, wall_source="code-ink-and-both-ais"
    )
    _run_current_checks(session, package_id, tmp_path)
    revision = session.scalars(
        select(PackageRevision).where(PackageRevision.package_id == package_id)
    ).one()
    live = session.scalars(
        select(Finding)
        .join(CheckRun, CheckRun.id == Finding.check_run_id)
        .where(Finding.package_revision_id == revision.id, CheckRun.superseded_at.is_(None))
    ).all()
    revision.state = "APPROVED"
    approval = Approval(package_revision_id=revision.id, approved_by="ana@example.com")
    session.add(approval)
    session.flush()
    session.add_all(
        ApprovedFinding(approval_id=approval.id, package_revision_id=revision.id, finding_id=f.id)
        for f in live
    )
    request_signed_exports(session, approval.id)
    session.commit()
    return anchors, revision, approval


def _change_a_wall_after_sign_off(session: Session, row_id: Any) -> None:
    """What a write after approval did before #1182; the API refuses it now, so straight in."""
    session.add(
        SlotRowReviewDecision(
            id=uuid4(),
            row_candidate_id=row_id,
            supersedes_id=None,
            wall_config="island",
            measurements={},
            confirmed_by="someone after the sign-off",
        )
    )
    session.commit()


def test_the_snapshot_freezes_the_countertop_results_the_screen_showed(
    session: Session, tmp_path: Path
) -> None:
    anchors, revision, approval = _approved(session, tmp_path)
    snapshot = session.scalars(
        select(ApprovalExportSnapshot).where(ApprovalExportSnapshot.approval_id == approval.id)
    ).one()
    screen_at_sign_off = _countertop_results_for_revision(session, revision.package_id, revision)

    stored = load_snapshot(session, snapshot).countertop_results
    assert stored is not None and stored.schema_version == 1
    frozen = printable_countertops(stored)

    assert frozen == screen_at_sign_off
    assert {item.row_id for item in frozen.items} >= set(anchors.values())


def test_a_wall_changed_after_sign_off_never_reaches_the_signed_files(
    session: Session, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Input: approve, then set row 1's wall to `island` in the database. Outcome: the live results
    show the change (so the test can fail), the signed workbook, PDF and redline labels still show
    the wall from the drawing that was approved."""
    anchors, revision, approval = _approved(session, tmp_path)
    store = LocalStore(root=tmp_path / "signed", ticket_secret=b"synthetic-test")
    stored = load_snapshot(
        session,
        session.scalars(
            select(ApprovalExportSnapshot).where(ApprovalExportSnapshot.approval_id == approval.id)
        ).one(),
    ).countertop_results
    assert stored is not None
    frozen = printable_countertops(stored)
    row = anchors[0]
    approved_wall = next(item.wall_layout for item in frozen.items if item.row_id == row)

    _change_a_wall_after_sign_off(session, row)
    live = _countertop_results_for_revision(session, revision.package_id, revision)
    live_wall = next(item.wall_layout for item in live.items if item.row_id == row)
    assert live_wall != approved_wall, "the change must be visible live, or this proves nothing"

    labels: list[tuple[str, ...]] = []

    def redline(_db: Session, the_store: LocalStore, **given: Any) -> RedlineOutput:
        """The drawing pages of this synthetic package were never stored; the labels it is handed
        are what matters here."""
        labels.append(tuple(given["countertop_labels"]))
        page = io.BytesIO()
        writer = PdfWriter()
        writer.add_blank_page(width=72, height=72)
        writer.write(page)
        page.seek(0)
        return RedlineOutput(
            the_store.put(f"redline/{uuid4()}.pdf", page, content_type="application/pdf"), None
        )

    monkeypatch.setattr("workflow.signed_outputs.render_evidence_grounded_redline", redline)
    bundle = generate_signed_outputs(session, store, approval.id)
    session.commit()

    workbook = session.get(OutputArtifact, bundle.workbook_id)
    pdf = session.get(OutputArtifact, bundle.pdf_id)
    assert workbook is not None and pdf is not None
    sheet = load_workbook(io.BytesIO(store.get(workbook.storage_key).read()))["Countertops"]
    values = {cell.value for line in sheet.iter_rows() for cell in line}
    assert approved_wall.label in values
    assert live_wall.label not in values
    text = " ".join(
        page.extract_text() or ""
        for page in PdfReader(io.BytesIO(store.get(pdf.storage_key).read())).pages
    )
    assert approved_wall.label is not None and live_wall.label is not None
    assert approved_wall.label in text and live_wall.label not in text
    assert labels == [tuple(f"Page {i.page_number} — {i.label}" for i in frozen.items)]


def test_a_snapshot_frozen_before_1182_still_exports_through_the_live_results(
    session: Session, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """D2. Input: an approval whose snapshot has no countertop section, as every snapshot made before
    #1182 (the key absent from its JSON). Outcome: it loads, and the signed files are drawn from the
    live results, exactly as before #1182."""
    import json

    from app.review.signed_exports import ExportSnapshot, frozen_countertop_results

    monkeypatch.setattr("app.review.signed_exports.freeze_countertops", lambda _results: None)
    anchors, revision, approval = _approved(session, tmp_path)
    snapshot = session.scalars(
        select(ApprovalExportSnapshot).where(ApprovalExportSnapshot.approval_id == approval.id)
    ).one()
    payload = load_snapshot(session, snapshot)
    assert payload.countertop_results is None
    old_json = json.loads(snapshot.canonical_json)
    del old_json["countertop_results"]
    assert ExportSnapshot.model_validate(old_json).countertop_results is None

    live = _countertop_results_for_revision(session, revision.package_id, revision)
    assert frozen_countertop_results(session, payload) == live

    store = LocalStore(root=tmp_path / "signed", ticket_secret=b"synthetic-test")

    def redline(_db: Session, the_store: LocalStore, **_given: Any) -> RedlineOutput:
        page = io.BytesIO()
        writer = PdfWriter()
        writer.add_blank_page(width=72, height=72)
        writer.write(page)
        page.seek(0)
        return RedlineOutput(
            the_store.put(f"redline/{uuid4()}.pdf", page, content_type="application/pdf"), None
        )

    monkeypatch.setattr("workflow.signed_outputs.render_evidence_grounded_redline", redline)
    bundle = generate_signed_outputs(session, store, approval.id)
    session.commit()

    workbook = session.get(OutputArtifact, bundle.workbook_id)
    assert workbook is not None
    sheet = load_workbook(io.BytesIO(store.get(workbook.storage_key).read()))["Countertops"]
    labels = {cell.value for line in sheet.iter_rows() for cell in line}
    assert {item.label for item in live.items} <= labels
    assert anchors
